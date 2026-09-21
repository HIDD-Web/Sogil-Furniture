from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, UploadFile, File
from starlette.middleware.cors import CORSMiddleware
from starlette.responses import Response as StarletteResponse
from motor.motor_asyncio import AsyncIOMotorClient
import logging
from pydantic import BaseModel
from typing import List, Optional, Dict, Any
import uuid
import json
import re
import urllib.parse
import jwt
import bcrypt
import requests
import boto3
from datetime import datetime, timezone, timedelta
from bson import ObjectId
from bson.errors import InvalidId

import certifi

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url, tlsCAFile=certifi.where())
db = client[os.environ['DB_NAME']]

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

app = FastAPI()
api_router = APIRouter(prefix="/api")
JWT_ALGORITHM = "HS256"

def get_jwt_secret() -> str:
    return os.environ["JWT_SECRET"]

def oid(v: str) -> ObjectId:
    try:
        return ObjectId(v)
    except (InvalidId, TypeError):
        raise HTTPException(status_code=404, detail="Data tidak ditemukan")

def id_query(v) -> dict:
    values = [v]

    if isinstance(v, str):
        try:
            values.append(ObjectId(v))
        except (InvalidId, TypeError):
            pass

    return {"$in": values}

# --------------------------------------------------------------------------
# Object storage
# --------------------------------------------------------------------------
R2_ENDPOINT = os.environ["R2_ENDPOINT"]
R2_ACCESS_KEY_ID = os.environ["R2_ACCESS_KEY_ID"]
R2_SECRET_ACCESS_KEY = os.environ["R2_SECRET_ACCESS_KEY"]
R2_BUCKET = os.environ.get("R2_BUCKET", "sogil-furniture")
APP_NAME = "sogil-furniture"
storage_key = None
MIME_TYPES = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "gif": "image/gif", "webp": "image/webp"}

def get_r2_client():
    return boto3.client(
        "s3",
        endpoint_url=R2_ENDPOINT,
        region_name="auto",
        aws_access_key_id=R2_ACCESS_KEY_ID,
        aws_secret_access_key=R2_SECRET_ACCESS_KEY,
    )

def put_object(path: str, data: bytes, content_type: str) -> dict:
    s3 = get_r2_client()

    s3.put_object(
        Bucket=R2_BUCKET,
        Key=path,
        Body=data,
        ContentType=content_type,
    )

    return {
        "storage_path": path,
        "size": len(data),
        "content_type": content_type,
    }

def get_object(path: str):
    s3 = get_r2_client()

    response = s3.get_object(
        Bucket=R2_BUCKET,
        Key=path,
    )

    content = response["Body"].read()
    content_type = response.get(
        "ContentType",
        "application/octet-stream",
    )

    return content, content_type

# --------------------------------------------------------------------------
# Auth & permissions
# --------------------------------------------------------------------------
ACCOUNT_TYPES = ["owner", "manager", "admin", "employee"]
PERMISSION_KEYS = ["manage_orders", "modify_products", "manage_settings", "access_finance", "delete_data", "manage_admins"]

def default_permissions(role: str) -> Dict[str, bool]:
    if role == "owner":
        return {k: True for k in PERMISSION_KEYS}
    if role == "manager":
        return {"manage_orders": True, "modify_products": True, "manage_settings": True, "access_finance": True, "delete_data": False, "manage_admins": False}
    if role == "admin":
        return {"manage_orders": True, "modify_products": True, "manage_settings": True, "access_finance": False, "delete_data": False, "manage_admins": False}
    return {"manage_orders": True, "modify_products": False, "manage_settings": False, "access_finance": False, "delete_data": False, "manage_admins": False}

def hash_password(p: str) -> str:
    return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()

def verify_password(p: str, h: str) -> bool:
    try:
        return bcrypt.checkpw(p.encode(), h.encode())
    except Exception:
        return False

def create_access_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email, "exp": datetime.now(timezone.utc) + timedelta(hours=12), "type": "access"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

async def get_current_admin(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        h = request.headers.get("Authorization", "")
        if h.startswith("Bearer "):
            token = h[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Token tidak valid")
        admin_id = oid(payload["sub"])

        user = await db.admins.find_one(
            {"_id": {"$in": [payload["sub"], admin_id]}}
        )

        if not user:
            raise HTTPException(status_code=401, detail="Pengguna tidak ditemukan")

        user["id"] = str(user.pop("_id"))
        user.pop("password_hash", None)
        if user.get("role") == "owner":
            user["permissions"] = default_permissions("owner")
        else:
            user.setdefault("permissions", default_permissions(user.get("role", "admin")))
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir, silakan login kembali")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")

def require_perm(perm: str):
    async def dep(admin: dict = Depends(get_current_admin)) -> dict:
        if admin.get("role") == "owner" or admin.get("permissions", {}).get(perm):
            return admin
        raise HTTPException(status_code=403, detail="Anda tidak memiliki akses untuk tindakan ini")
    return dep

def require_owner():
    async def dep(admin: dict = Depends(get_current_admin)) -> dict:
        if admin.get("role") != "owner":
            raise HTTPException(status_code=403, detail="Hanya Owner/CEO yang dapat melakukan ini")
        return admin
    return dep

def create_customer_token(cid: str) -> str:
    payload = {"sub": cid, "exp": datetime.now(timezone.utc) + timedelta(days=30), "type": "customer"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

async def get_optional_customer(request: Request):
    token = request.cookies.get("customer_token")
    if not token:
        h = request.headers.get("X-Customer-Authorization", "")
        if h.startswith("Bearer "):
            token = h[7:]
    if not token:
        return None
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "customer":
            return None
        c = await db.customers.find_one({"_id": id_query(payload["sub"])})
        if c and c.get("active") is False:
            return None
        return c
    except jwt.InvalidTokenError:
        return None

async def get_current_customer(request: Request):
    c = await get_optional_customer(request)
    if not c:
        raise HTTPException(status_code=401, detail="Silakan login sebagai pelanggan")
    return c

def audit_fields(admin: dict) -> dict:
    return {"updated_at": datetime.now(timezone.utc).isoformat(),
            "updated_by_id": admin.get("id"), "updated_by_name": admin.get("name", "Admin")}

# --------------------------------------------------------------------------
# Models
# --------------------------------------------------------------------------
class LoginInput(BaseModel):
    email: str
    password: str

class ChangePasswordInput(BaseModel):
    current_password: str
    new_password: str

class ChangeEmailInput(BaseModel):
    current_password: str
    new_email: str

class DeliveryZoneInput(BaseModel):
    name: str
    fee_le: float
    active: bool = True

class OrderItemInput(BaseModel):
    product_id: str
    config: Dict[str, Any] = {}
    quantity: int = 1

class OrderInput(BaseModel):
    customer_name: str
    customer_phone: str
    customer_address: str
    customer_maps_url: Optional[str] = ""
    delivery_method: str
    delivery_zone_id: Optional[str] = None
    payment_method: str
    notes: Optional[str] = ""
    item: Optional[OrderItemInput] = None
    items: Optional[List[OrderItemInput]] = None
    discount_code: Optional[str] = None
    referral_code: Optional[str] = None
    redeem_points: Optional[float] = 0

class OrderStatusUpdate(BaseModel):
    order_status: Optional[str] = None
    payment_status: Optional[str] = None
    admin_note: Optional[str] = None
    subtotal_le: Optional[float] = None
    delivery_fee_le: Optional[float] = None
    discount_le: Optional[float] = None
    total_le: Optional[float] = None
    items: Optional[List[dict]] = None

class CustomRequestInput(BaseModel):
    customer_name: str
    customer_phone: str
    customer_address: Optional[str] = ""
    furniture_type: str
    dimensions: Optional[Dict[str, Any]] = None
    material: Optional[str] = ""
    reference_photos: Optional[List[str]] = []
    budget_estimation_le: Optional[float] = None
    notes: Optional[str] = ""

class CustomRequestUpdate(BaseModel):
    status: Optional[str] = None
    admin_notes: Optional[str] = None
    estimated_price_le: Optional[float] = None

class ConvertCustomToOrderInput(BaseModel):
    subtotal_le: float
    delivery_fee_le: Optional[float] = 0.0
    delivery_method: str = "pickup"
    delivery_zone_id: Optional[str] = None
    payment_method: str = "transfer"
    customer_address: Optional[str] = None
    customer_maps_url: Optional[str] = None
    notes: Optional[str] = None

class AdminCreateInput(BaseModel):
    name: str
    email: str
    password: str
    role: str = "admin"
    permissions: Optional[Dict[str, bool]] = None

class AdminUpdateInput(BaseModel):
    name: Optional[str] = None
    email: Optional[str] = None
    password: Optional[str] = None
    role: Optional[str] = None
    permissions: Optional[Dict[str, bool]] = None
    status: Optional[str] = None

class PublishCustomCollectionInput(BaseModel):
    title: str
    price_le: float
    spesifikasi: Optional[str] = ""
    photo_urls: List[str] = []
    order_id: Optional[str] = None
    request_id: Optional[str] = None

class FinanceTxnInput(BaseModel):
    date: Optional[str] = None
    type: str  # income | expense
    category: str
    amount: float
    currency: str  # IDR | EGP
    description: Optional[str] = ""
    recipient_employee_id: Optional[str] = None
    exchange_rate: Optional[float] = None
    counterpart_amount: Optional[float] = None

class TransferInput(BaseModel):
    date: Optional[str] = None
    from_account: str
    to_account: str
    from_amount: float
    to_amount: float
    exchange_rate: Optional[float] = None
    description: Optional[str] = ""

class CustomerRegister(BaseModel):
    username: str
    phone: str
    password: str
    email: Optional[str] = ""

class CustomerLogin(BaseModel):
    identifier: str  # phone or username
    password: str

class ReferralInput(BaseModel):
    customer_id: str
    code: str
    status: str = "active"
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    discount_percentage: float = 0
    max_discount_le: float = 0
    max_claim_orders: int = 0
    max_reward_orders: int = 0
    max_total_points: int = 0
    points_per_order: float = 0

class PointAdjustInput(BaseModel):
    amount: float
    reason: str

# --------------------------------------------------------------------------
# Pricing engine
# --------------------------------------------------------------------------
def _num(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default

def compute_item_price(product: dict, config: dict, quantity: int):
    pricing = product.get("pricing", {}) or {}
    category = product.get("category")
    quantity = max(1, int(quantity or 1))
    bd = {"base_price_le": 0.0, "adjustments_le": 0.0, "finishing_le": 0.0, "unit_price_le": 0.0,
          "subtotal_le": 0.0, "lines": [], "requires_admin_confirmation": False, "label": product.get("name", "")}

    if bool(config.get("custom_size")):
        bd["requires_admin_confirmation"] = True
        bd["lines"].append({"label": "Ukuran Custom", "value": 0.0, "note": "Akan dikonfirmasi admin"})
        return bd

    if pricing.get("price_model") == "additive":
        base_keys = pricing.get("base_keys", []) or []
        base_prices = pricing.get("base_prices", {}) or {}
        groups = {g.get("key"): g for g in (pricing.get("groups") or [])}
        if base_keys:
            key = " | ".join(str(config.get(k, "")) for k in base_keys)
            base = _num(base_prices.get(key))
            bd["base_price_le"] = base
            label_vals = ", ".join(str(config.get(k, "")) for k in base_keys if config.get(k))
            bd["lines"].append({"label": label_vals or "Harga dasar", "value": base})
        for gkey, mapping in (pricing.get("adjust", {}) or {}).items():
            val = config.get(gkey)
            amt = _num((mapping or {}).get(val))
            if amt:
                bd["adjustments_le"] += amt
                glabel = (groups.get(gkey) or {}).get("label", gkey)
                bd["lines"].append({"label": f"{glabel}: {val}", "value": amt})
        unit = bd["base_price_le"] + bd["adjustments_le"] + bd["finishing_le"]
        bd["unit_price_le"] = unit
        bd["subtotal_le"] = unit * quantity
        return bd

    if category == "rak":
        length = str(config.get("length", "")); level = str(config.get("level", ""))
        rtype = config.get("type", "B"); finishing = config.get("finishing", "Natural")
        base = _num((pricing.get("base_prices", {}) or {}).get(f"{length}_{level}"))
        bd["base_price_le"] = base
        bd["lines"].append({"label": f"Harga dasar ({length} cm, {level} tingkat)", "value": base})
        type_adj = _num(((pricing.get("type_adjustments", {}) or {}).get(rtype, {}) or {}).get(length))
        if type_adj:
            bd["adjustments_le"] += type_adj
            bd["lines"].append({"label": f"Tipe {rtype}", "value": type_adj})
        fin = (pricing.get("finishing", {}) or {}).get(finishing, 0)
        fin_cost = _num(fin.get(length)) * _num(level, 0) if isinstance(fin, dict) else _num(fin)
        if fin_cost:
            bd["finishing_le"] = fin_cost
            bd["lines"].append({"label": f"Finishing {finishing}", "value": fin_cost})
    elif category == "meja":
        size = str(config.get("size", "")); height = str(config.get("height", "")); finishing = config.get("finishing", "Natural")
        base = _num((pricing.get("base_prices", {}) or {}).get(f"{size}_{height}"))
        bd["base_price_le"] = base
        bd["lines"].append({"label": f"Harga dasar ({size} cm, tinggi {height} cm)", "value": base})
        fin = (pricing.get("finishing", {}) or {}).get(finishing, 0)
        fin_cost = _num(fin.get(f"{size}_{height}")) if isinstance(fin, dict) else _num(fin)
        if fin_cost:
            bd["finishing_le"] = fin_cost
            bd["lines"].append({"label": f"Finishing {finishing}", "value": fin_cost})
    elif category == "meja_rak":
        variant = str(config.get("variant", "")); rtype = config.get("type", "B"); finishing = config.get("finishing", "Natural")
        base = _num((pricing.get("base_prices", {}) or {}).get(variant))
        bd["base_price_le"] = base
        bd["lines"].append({"label": f"Harga dasar ({variant})", "value": base})
        tam = (pricing.get("type_adjustments", {}) or {}).get(rtype, {}) or {}
        type_adj = _num(tam.get(variant)) if isinstance(tam, dict) else 0.0
        if type_adj:
            bd["adjustments_le"] += type_adj
            bd["lines"].append({"label": f"Tipe {rtype}", "value": type_adj})
        fin = (pricing.get("finishing", {}) or {}).get(finishing, 0)
        fin_cost = _num(fin) if not isinstance(fin, dict) else 0.0
        if fin_cost:
            bd["finishing_le"] = fin_cost
            bd["lines"].append({"label": f"Finishing {finishing}", "value": fin_cost})
    else:
        bd["requires_admin_confirmation"] = True

    unit = bd["base_price_le"] + bd["adjustments_le"] + bd["finishing_le"]
    bd["unit_price_le"] = unit
    bd["subtotal_le"] = unit * quantity
    return bd

# --------------------------------------------------------------------------
# Photo weighted similarity
# --------------------------------------------------------------------------
DEFAULT_WEIGHTS = {
    "rak": {"length": 0.45, "level": 0.35, "type": 0.10, "finishing": 0.10},
    "meja": {"size": 0.5, "height": 0.3, "finishing": 0.2},
    "meja_rak": {"variant": 0.6, "type": 0.2, "finishing": 0.2},
}

def _numeric_sim(a, b, scale):
    try:
        return max(0.0, 1.0 - abs(float(a) - float(b)) / scale)
    except (TypeError, ValueError):
        return 1.0 if str(a) == str(b) else 0.0

def photo_similarity(category, weights, req, attr):
    score = 0.0
    for key, w in weights.items():
        rv, av = req.get(key), attr.get(key)
        if rv in (None, "") or av in (None, ""):
            continue
        if key in ("length", "level", "height"):
            scale = {"length": 60.0, "level": 5.0, "height": 45.0}[key]
            score += w * _numeric_sim(rv, av, scale)
        else:
            score += w * (1.0 if str(rv) == str(av) else 0.0)
    return score

def match_photo(product, config):
    photos = product.get("photos", []) or []
    if not photos:
        return {"photo": None, "exact": False}
    category = product.get("category")
    weights = (product.get("photo_weights") or {}) or DEFAULT_WEIGHTS.get(category, {})
    keys = list(weights.keys())
    best, best_score, best_exact = None, -1.0, False
    for ph in photos:
        attr = ph.get("attributes", {}) or {}
        exact = all(str(config.get(k, "")) == str(attr.get(k, "")) for k in keys if attr.get(k) not in (None, ""))
        s = photo_similarity(category, weights, config, attr)
        if exact:
            s += 100
        if s > best_score:
            best, best_score, best_exact = ph, s, exact
    return {"photo": best, "exact": best_exact}

# --------------------------------------------------------------------------
# Settings helpers
# --------------------------------------------------------------------------
async def get_setting(key, default=None):
    doc = await db.settings.find_one({"key": key})
    return doc["value"] if doc else default

async def set_setting(key, value, admin=None):
    upd = {"key": key, "value": value, "updated_at": datetime.now(timezone.utc).isoformat()}
    if admin:
        upd["updated_by_name"] = admin.get("name", "Admin")
    await db.settings.update_one({"key": key}, {"$set": upd}, upsert=True)

def clean(doc):
    if not doc:
        return doc
    doc = dict(doc)
    if "_id" in doc:
        doc["id"] = str(doc.pop("_id"))
    doc.pop("password_hash", None)
    return doc

# --------------------------------------------------------------------------
# Auth routes
# --------------------------------------------------------------------------
@api_router.post("/auth/login")
async def login(data: LoginInput, response: Response):
    email = data.email.strip().lower()
    user = await db.admins.find_one({"email": email})
    if not user or not verify_password(data.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Email atau kata sandi salah")
    token = create_access_token(str(user["_id"]), email)
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=43200, path="/")
    return {"id": str(user["_id"]), "email": email, "name": user.get("name", "Admin"),
            "role": user.get("role", "admin"),
            "token": token,
            "permissions": default_permissions("owner") if user.get("role") == "owner" else user.get("permissions", default_permissions(user.get("role", "admin")))}

@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}

@api_router.get("/auth/me")
async def me(admin: dict = Depends(get_current_admin)):
    return admin

@api_router.post("/auth/change-password")
async def change_password(data: ChangePasswordInput, admin: dict = Depends(get_current_admin)):
    user = await db.admins.find_one({"_id": id_query(admin["id"])})
    if not verify_password(data.current_password, user["password_hash"]):
        raise HTTPException(status_code=400, detail="Kata sandi saat ini salah")
    if len(data.new_password) < 6:
        raise HTTPException(status_code=400, detail="Kata sandi baru minimal 6 karakter")
    await db.admins.update_one({"_id": id_query(admin["id"])}, {"$set": {"password_hash": hash_password(data.new_password)}})
    return {"ok": True}

@api_router.post("/auth/change-email")
async def change_email(data: ChangeEmailInput, response: Response, admin: dict = Depends(get_current_admin)):
    user = await db.admins.find_one({"_id": id_query(admin["id"])})
    if not verify_password(data.current_password, user["password_hash"]):
        raise HTTPException(status_code=400, detail="Kata sandi salah")
    new_email = data.new_email.strip().lower()
    if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", new_email):
        raise HTTPException(status_code=400, detail="Format email tidak valid")
    if await db.admins.find_one({"email": new_email, "_id": {"$ne": oid(admin["id"])}}):
        raise HTTPException(status_code=400, detail="Email sudah digunakan")
    await db.admins.update_one({"_id": id_query(admin["id"])}, {"$set": {"email": new_email}})
    token = create_access_token(admin["id"], new_email)
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=43200, path="/")
    return {"ok": True, "email": new_email}

# --------------------------------------------------------------------------
# Public product routes
# --------------------------------------------------------------------------
async def auto_category_cover(product):
    """Most frequently ordered valid configuration's photo for this category."""
    category = product.get("category")
    agg = {}
    async for o in db.orders.find({"order_status": {"$ne": "dibatalkan"}}):
        for it in o.get("items", [o.get("item")]):
            if not it or it.get("category") != category:
                continue
            cfg = it.get("configuration_snapshot", {}) or {}
            key = tuple(sorted((k, str(v)) for k, v in cfg.items() if k != "custom_note"))
            agg[key] = agg.get(key, {"count": 0, "cfg": cfg})
            agg[key]["count"] += int(it.get("quantity", 1))
    for _, e in sorted(agg.items(), key=lambda x: x[1]["count"], reverse=True):
        m = match_photo(product, e["cfg"])
        if m.get("photo"):
            ph = m["photo"]
            url = ph.get("main_url") or ph.get("front_url") or ph.get("side_url")
            if url:
                return url
    return ""

async def auto_category_cover_detail(product):
    """Preview details for automatic category cover: winning config, order count, matched image."""
    category = product.get("category")
    agg = {}
    async for o in db.orders.find({"order_status": {"$ne": "dibatalkan"}}):
        for it in o.get("items", [o.get("item")]):
            if not it or it.get("category") != category:
                continue
            cfg = it.get("configuration_snapshot", {}) or {}
            key = tuple(sorted((k, str(v)) for k, v in cfg.items() if k != "custom_note"))
            agg[key] = agg.get(key, {"count": 0, "cfg": cfg})
            agg[key]["count"] += int(it.get("quantity", 1))
    ordered = sorted(agg.items(), key=lambda x: x[1]["count"], reverse=True)
    for _, e in ordered:
        m = match_photo(product, e["cfg"])
        ph = m.get("photo")
        if ph:
            url = ph.get("main_url") or ph.get("front_url") or ph.get("side_url")
            if url:
                return {"config": e["cfg"], "count": e["count"], "image": url}
    if ordered:
        return {"config": ordered[0][1]["cfg"], "count": ordered[0][1]["count"], "image": ""}
    return None

async def resolve_display_photo(product):
    """Category COVER image only. Never falls back to a specific config photo (except in auto mode)."""
    if product.get("cover_mode") == "auto":
        auto = await auto_category_cover(product)
        if auto:
            return auto
    return product.get("category_cover_image") or product.get("image_url") or ""

# --------------------------------------------------------------------------
# Categories (Dynamic Categories)
# --------------------------------------------------------------------------
@api_router.get("/categories")
async def list_categories():
    cats = await db.categories.find({"active": True}).sort("sort_order", 1).to_list(100)
    return [clean(c) for c in cats]

@api_router.get("/admin/categories")
async def list_admin_categories(admin: dict = Depends(require_perm("modify_products"))):
    cats = await db.categories.find({}).sort("sort_order", 1).to_list(100)
    return [clean(c) for c in cats]

@api_router.post("/admin/categories")
async def create_category(payload: dict, admin: dict = Depends(require_perm("modify_products"))):
    key = str(payload.get("key", "")).strip().lower().replace(" ", "_")
    name = str(payload.get("name", "")).strip()
    if not key or not name:
        raise HTTPException(status_code=400, detail="Key dan Nama kategori wajib diisi")
    existing = await db.categories.find_one({"key": key})
    if existing:
        raise HTTPException(status_code=400, detail=f"Kategori dengan key '{key}' sudah ada")
    count = await db.categories.count_documents({})
    doc = {
        "key": key,
        "name": name,
        "description": str(payload.get("description", "")).strip(),
        "sort_order": int(payload.get("sort_order", count + 1)),
        "active": bool(payload.get("active", True)),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    res = await db.categories.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    return clean(doc)

@api_router.put("/admin/categories/{category_id}")
async def update_category(category_id: str, payload: dict, admin: dict = Depends(require_perm("modify_products"))):
    cat = await db.categories.find_one({"_id": id_query(category_id)})
    if not cat:
        raise HTTPException(status_code=404, detail="Kategori tidak ditemukan")
    upd = {"updated_at": datetime.now(timezone.utc).isoformat()}
    for field in ["name", "description", "sort_order", "active", "featured_preview_ids", "price_list_image_url"]:
        if field in payload:
            upd[field] = payload[field]
    if "key" in payload:
        new_key = str(payload["key"]).strip().lower().replace(" ", "_")
        if new_key and new_key != cat.get("key"):
            existing = await db.categories.find_one({"key": new_key, "_id": {"$ne": cat["_id"]}})
            if existing:
                raise HTTPException(status_code=400, detail=f"Key '{new_key}' sudah digunakan")
            upd["key"] = new_key
    await db.categories.update_one({"_id": cat["_id"]}, {"$set": upd})
    updated = await db.categories.find_one({"_id": cat["_id"]})
    return clean(updated)

@api_router.delete("/admin/categories/{category_id}")
async def delete_category(category_id: str, admin: dict = Depends(require_perm("modify_products"))):
    cat = await db.categories.find_one({"_id": id_query(category_id)})
    if not cat:
        raise HTTPException(status_code=404, detail="Kategori tidak ditemukan")
    has_products = await db.products.count_documents({"category": cat.get("key")})
    if has_products > 0:
        await db.categories.update_one({"_id": cat["_id"]}, {"$set": {"active": False, "updated_at": datetime.now(timezone.utc).isoformat()}})
        return {"status": "deactivated", "message": "Kategori dinonaktifkan"}
    await db.categories.delete_one({"_id": cat["_id"]})
    return {"status": "deleted"}

@api_router.get("/products")
async def list_products(admin_view: bool = False):
    query = {} if admin_view else {"active": True}
    products = await db.products.find(query).sort("sort_order", 1).to_list(200)
    out = []
    for p in products:
        c = clean(p)
        c["display_image"] = await resolve_display_photo(c)
        out.append(c)
    return out

@api_router.get("/admin/products/{product_id}/cover-preview")
async def admin_cover_preview(product_id: str, admin: dict = Depends(get_current_admin)):
    product = await db.products.find_one({"_id": id_query(product_id)})
    if not product:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    product = clean(product)
    mode = product.get("cover_mode", "manual")
    selected = await auto_category_cover_detail(product) if mode == "auto" else None
    return {"cover_mode": mode, "category_cover_image": product.get("category_cover_image", ""),
            "resolved_image": await resolve_display_photo(product), "selected": selected}

@api_router.get("/products/{slug}")
async def get_product(slug: str):
    product = await db.products.find_one({"slug": slug})
    if not product:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    c = clean(product)
    c["display_image"] = await resolve_display_photo(c)
    return c

@api_router.post("/calculate-price")
async def calculate_price(data: OrderItemInput):
    product = await db.products.find_one({"_id": id_query(data.product_id)})
    if not product:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return compute_item_price(product, data.config, data.quantity)

@api_router.post("/match-photo")
async def match_photo_endpoint(data: OrderItemInput):
    product = await db.products.find_one({"_id": id_query(data.product_id)})
    if not product:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return match_photo(product, data.config)

# --------------------------------------------------------------------------
# Public settings & zones
# --------------------------------------------------------------------------
@api_router.get("/delivery-zones")
async def list_zones(admin_view: bool = False):
    query = {} if admin_view else {"active": True}
    zones = await db.delivery_zones.find(query).sort("fee_le", 1).to_list(100)
    return [clean(z) for z in zones]

@api_router.get("/store-info")
async def store_info():
    return {
        "store_name": await get_setting("store_name", "Sogil Furniture"),
        "app_name": await get_setting("app_name", "Sogil Furniture — Furniture Ordering & Management"),
        "tagline": await get_setting("tagline", "Kualitas Terbaik, Untuk Ruang Terbaik"),
        "store_address": await get_setting("store_address", ""),
        "store_maps_url": await get_setting("store_maps_url", ""),
        "whatsapp_number": await get_setting("whatsapp_number", "628XXXXXXXXXX"),
        "exchange_rate_idr_per_le": _num(await get_setting("exchange_rate_idr_per_le", 357), 357),
        "bank_info": await get_setting("bank_info", ""),
        "logo_url": await get_setting("logo_url", ""),
        "referral_program_enabled": await get_setting("referral_program_enabled", True),
        "point_redeem_max_pct": _num(await get_setting("point_redeem_max_pct", 50), 50),
        "instagram": await get_setting("instagram", ""),
        "facebook": await get_setting("facebook", ""),
        "tiktok": await get_setting("tiktok", ""),
        "email": await get_setting("email", ""),
    }

# --------------------------------------------------------------------------
# Orders
# --------------------------------------------------------------------------
def normalize_phone(raw: str) -> str:
    s = re.sub(r"[\s\-().]", "", raw or "")
    return s

def valid_intl_phone(s: str) -> bool:
    return bool(re.match(r"^\+\d{8,15}$", s))

async def generate_order_number():
    today = datetime.now(timezone.utc).strftime("%Y%m%d")
    prefix = f"SGF-{today}-"
    count = await db.orders.count_documents({"order_number": {"$regex": f"^{prefix}"}})
    seq = count + 1
    while await db.orders.find_one({"order_number": f"{prefix}{seq:03d}"}):
        seq += 1
    return f"{prefix}{seq:03d}"

@api_router.post("/orders")
async def create_order(request: Request, data: OrderInput):
    if not data.customer_name.strip():
        raise HTTPException(status_code=400, detail="Nama tidak boleh kosong.")
    phone = normalize_phone(data.customer_phone)
    if not phone:
        raise HTTPException(status_code=400, detail="No HP tidak boleh kosong.")
    if not valid_intl_phone(phone):
        raise HTTPException(status_code=400, detail="No HP harus diawali + dan kode negara. Contoh: +62xxxxxxxxxx")
    if not data.customer_address.strip():
        raise HTTPException(status_code=400, detail="Alamat tidak boleh kosong.")
    if data.delivery_method not in ("delivery", "pickup"):
        raise HTTPException(status_code=400, detail="Silakan pilih metode penerimaan barang.")
    if data.payment_method not in ("cash", "transfer"):
        raise HTTPException(status_code=400, detail="Silakan pilih metode pembayaran.")

    items_input = data.items if data.items else ([data.item] if data.item else [])
    if not items_input:
        raise HTTPException(status_code=400, detail="Keranjang kosong.")

    maps = (data.customer_maps_url or "").strip()
    if maps and not re.match(r"^https?://", maps):
        raise HTTPException(status_code=400, detail="Link Maps tidak valid.")

    delivery_fee = 0.0; zone_name = None; zone_id = None
    if data.delivery_method == "delivery":
        if not data.delivery_zone_id:
            raise HTTPException(status_code=400, detail="Silakan pilih zona pengiriman.")
        zone = await db.delivery_zones.find_one({"_id": id_query(data.delivery_zone_id)})
        if not zone or not zone.get("active"):
            raise HTTPException(status_code=400, detail="Zona pengiriman tidak valid.")
        delivery_fee = _num(zone["fee_le"]); zone_name = zone["name"]; zone_id = str(zone["_id"])

    items = []; subtotal_le = 0.0; requires_confirm = False
    for it in items_input:
        if it.quantity < 1:
            raise HTTPException(status_code=400, detail="Jumlah minimal 1.")
        product = await db.products.find_one({"_id": id_query(it.product_id)})
        if not product or not product.get("active"):
            raise HTTPException(status_code=400, detail="Produk tidak tersedia.")
        bd = compute_item_price(product, it.config, it.quantity)
        subtotal_le += bd["subtotal_le"]
        requires_confirm = requires_confirm or bd["requires_admin_confirmation"] or bool(it.config.get("custom_size"))
        items.append({
            "product_id": str(product["_id"]), "product_name_snapshot": product["name"], "category": product["category"],
            "configuration_snapshot": it.config, "quantity": it.quantity,
            "base_price_le": bd["base_price_le"], "adjustments_le": bd["adjustments_le"],
            "finishing_le": bd["finishing_le"], "unit_price_le": bd["unit_price_le"],
            "subtotal_le": bd["subtotal_le"], "lines": bd["lines"],
        })

    total_le = subtotal_le + delivery_fee
    buyer = await get_optional_customer(request)
    buyer_id = str(buyer["_id"]) if buyer else None
    # Referral OR promo discount (exclusive to protect margin)
    referral_snap = None
    discount_le = 0.0
    disc_doc = None
    if data.referral_code:
        if not buyer:
            raise HTTPException(status_code=401, detail="Login sebagai pelanggan untuk memakai kode referral")
        rdisc, rdoc, rerr = await validate_referral(data.referral_code, subtotal_le, buyer_id)
        if rerr:
            raise HTTPException(status_code=400, detail=rerr)
        discount_le = rdisc
        owner = await db.customers.find_one({"_id": id_query(rdoc["customer_id"])})
        referral_snap = {"code": rdoc.get("code"), "owner_id": str(rdoc["customer_id"]),
                         "owner_name": owner.get("username") if owner else "", "percentage": _num(rdoc.get("discount_percentage")),
                         "discount_le": rdisc, "points_per_order": _num(rdoc.get("points_per_order"))}
    else:
        discount_le, disc_doc, disc_err = await validate_discount(data.discount_code, subtotal_le)
        if data.discount_code and disc_err:
            raise HTTPException(status_code=400, detail=disc_err)
    # Points redemption (logged-in only, max % of subtotal)
    points_redeemed = 0.0
    redeem_req = _num(data.redeem_points)
    if redeem_req > 0:
        if not buyer:
            raise HTTPException(status_code=401, detail="Login untuk menukar poin")
        max_pct = _num(await get_setting("point_redeem_max_pct", 50), 50)
        cap = min(subtotal_le * max_pct / 100.0, _num(buyer.get("points_available")))
        points_redeemed = round(min(redeem_req, cap), 2)
        if points_redeemed < 0:
            points_redeemed = 0.0
    total_le = max(0.0, subtotal_le - discount_le - points_redeemed + delivery_fee)
    rate = _num(await get_setting("exchange_rate_idr_per_le", 357), 357)
    now = datetime.now(timezone.utc).isoformat()

    order_doc = {
        "order_number": await generate_order_number(),
        "customer_name": data.customer_name.strip(), "customer_phone": phone,
        "customer_address": data.customer_address.strip(), "customer_maps_url": maps,
        "delivery_method": data.delivery_method, "delivery_zone_id": zone_id, "delivery_zone_name": zone_name,
        "delivery_fee_le": delivery_fee, "payment_method": data.payment_method,
        "payment_status": "belum_dibayar", "order_status": "pesanan_masuk",
        "notes": (data.notes or "").strip(), "admin_note": "",
        "items": items, "item": items[0],
        "subtotal_le": subtotal_le, "discount_code": (disc_doc.get("code") if disc_doc else None),
        "discount_percentage": (_num(disc_doc.get("percentage")) if disc_doc else 0), "discount_le": (discount_le if not referral_snap else 0),
        "referral": referral_snap, "referral_discount_le": (discount_le if referral_snap else 0),
        "points_redeemed_le": points_redeemed, "customer_id": buyer_id,
        "customer_username": (buyer.get("username") if buyer else None),
        "total_le": total_le,
        "exchange_rate_idr_per_le": rate, "estimated_total_idr": round(total_le * rate),
        "rate_timestamp": now, "requires_admin_confirmation": requires_confirm,
        "created_at": now, "updated_at": now,
    }
    result = await db.orders.insert_one(order_doc)
    if disc_doc:
        await db.discounts.update_one({"_id": disc_doc["_id"]}, {"$inc": {"claims": 1}})
    if referral_snap:
        await db.referrals.update_one({"code": referral_snap["code"]}, {"$inc": {"claims": 1}})
    if points_redeemed > 0 and buyer_id:
        newb = _num(buyer.get("points_available")) - points_redeemed
        await db.customers.update_one({"_id": id_query(buyer_id)}, {"$inc": {"points_available": -points_redeemed, "points_redeemed": points_redeemed}})
        await db.point_transactions.insert_one({"customer_id": buyer_id, "type": "redeem", "amount": -points_redeemed,
            "order_id": str(result.inserted_id), "reason": "Penukaran poin", "balance_after": newb,
            "created_at": datetime.now(timezone.utc).isoformat()})
    order_doc["id"] = str(result.inserted_id); order_doc.pop("_id", None)
    order_doc["whatsapp_message"] = build_whatsapp_message(order_doc)
    order_doc["whatsapp_number"] = await get_setting("whatsapp_number", "628XXXXXXXXXX")
    return order_doc

def fmt_le(v):
    return f"{v:,.0f}".replace(",", ".")

def fmt_idr(v):
    return f"Rp{v:,.0f}".replace(",", ".")

def config_summary(item):
    cfg = item.get("configuration_snapshot", {}) or {}; cat = item.get("category"); parts = []
    if cfg.get("_summary"):
        return str(cfg["_summary"])
    if cat == "rak":
        if cfg.get("length"): parts.append(f"{cfg['length']} cm")
        if cfg.get("level"): parts.append(f"{cfg['level']} Tingkat")
        if cfg.get("type"): parts.append(f"Tipe {cfg['type']}")
        if cfg.get("finishing"): parts.append(cfg["finishing"])
    elif cat == "meja":
        if cfg.get("size"): parts.append(f"{cfg['size']} cm")
        if cfg.get("height"): parts.append(f"Tinggi {cfg['height']} cm")
        if cfg.get("finishing"): parts.append(cfg["finishing"])
    elif cat == "meja_rak":
        if cfg.get("variant"): parts.append(cfg["variant"])
        if cfg.get("type"): parts.append(f"Tipe {cfg['type']}")
        if cfg.get("finishing"): parts.append(cfg["finishing"])
    if cfg.get("custom_size"): parts.append("Ukuran Custom")
    return ", ".join(parts)

def build_whatsapp_message(order):
    lines = ["Halo Sogil Furniture, saya ingin melakukan pemesanan.", "", f"No. Pesanan: {order['order_number']}", ""]
    for idx, item in enumerate(order.get("items", [order.get("item")]), 1):
        if not item:
            continue
        prefix = f"{idx}. " if len(order.get("items", [])) > 1 else ""
        lines.append(f"{prefix}Produk: {item['product_name_snapshot']}")
        cs = config_summary(item)
        if cs:
            lines.append(f"   Spesifikasi: {cs}")
        lines.append(f"   Jumlah: {item['quantity']}")
    lines.append("")
    if order["delivery_method"] == "delivery":
        lines.append("Pengiriman: Delivery")
        lines.append(f"Zona: {order.get('delivery_zone_name', '-')}")
    else:
        lines.append("Pengiriman: Ambil di Toko")
    lines += ["", "Rincian Harga:", f"Subtotal: {fmt_le(_num(order.get('subtotal_le')))} LE"]
    if _num(order.get("discount_le")) > 0:
        _dc = order.get("discount_code")
        lines.append(f"Diskon{f' ({_dc})' if _dc else ''}: -{fmt_le(_num(order['discount_le']))} LE")
    if _num(order.get("referral_discount_le")) > 0:
        _rc = (order.get("referral") or {}).get("code")
        lines.append(f"Diskon Referral{f' ({_rc})' if _rc else ''}: -{fmt_le(_num(order['referral_discount_le']))} LE")
    if _num(order.get("points_redeemed_le")) > 0:
        lines.append(f"Penukaran Poin: -{fmt_le(_num(order['points_redeemed_le']))} LE")
    lines += [f"Ongkir: {fmt_le(_num(order.get('delivery_fee_le')))} LE",
              f"Total: {fmt_le(order['total_le'])} LE (≈ {fmt_idr(order['estimated_total_idr'])})",
              f"Rate: Rp{fmt_le(order['exchange_rate_idr_per_le'])}/LE", "",
              f"Nama: {order['customer_name']}", f"No. HP: {order['customer_phone']}", f"Alamat: {order['customer_address']}"]
    if order.get("customer_maps_url"):
        lines.append(f"Maps: {order['customer_maps_url']}")
    lines += ["", f"Pembayaran: {'Transfer' if order['payment_method'] == 'transfer' else 'Cash'}"]
    if order.get("notes"):
        lines += ["", f"Catatan: {order['notes']}"]
    if order.get("requires_admin_confirmation"):
        lines += ["", "(*Pesanan ini memerlukan konfirmasi admin)"]
    return "\n".join(lines)

@api_router.get("/orders/{order_id}")
async def get_order_public(order_id: str, request: Request):
    order = await db.orders.find_one({"_id": id_query(order_id)})
    if not order:
        raise HTTPException(status_code=404, detail="Pesanan tidak ditemukan")
    # Orders tied to a customer account are private to that customer.
    # Guest orders (no customer_id) stay accessible by ID for checkout confirmation.
    # Admins/Owner use the RBAC-gated /admin/orders endpoints.
    if order.get("customer_id"):
        buyer = await get_optional_customer(request)
        if not buyer or str(buyer["_id"]) != str(order["customer_id"]):
            raise HTTPException(status_code=403, detail="Tidak diizinkan mengakses pesanan ini")
    order = clean(order)
    order["whatsapp_message"] = build_whatsapp_message(order)
    order["whatsapp_number"] = await get_setting("whatsapp_number", "628XXXXXXXXXX")
    return order

# --------------------------------------------------------------------------
# Admin: orders
# --------------------------------------------------------------------------
@api_router.get("/admin/orders")
async def admin_list_orders(status: Optional[str] = None, payment_status: Optional[str] = None,
                            q: Optional[str] = None, admin: dict = Depends(require_perm("manage_orders"))):
    query = {}
    if status:
        query["order_status"] = status
    if payment_status:
        query["payment_status"] = payment_status
    if q and q.strip():
        rx = {"$regex": re.escape(q.strip()), "$options": "i"}
        query["$or"] = [{"order_number": rx}, {"customer_name": rx}, {"customer_phone": rx},
                        {"customer_username": rx}, {"items.product_name_snapshot": rx}, {"referral.code": rx}]
    orders = await db.orders.find(query).sort("created_at", -1).to_list(2000)
    return [clean(o) for o in orders]

@api_router.get("/admin/overview")
async def admin_overview(admin: dict = Depends(require_perm("manage_orders"))):
    orders = await db.orders.find({}).to_list(5000)
    stats = {"total": len(orders), "pesanan_masuk": 0, "dikonfirmasi": 0, "diproses": 0, "siap": 0,
             "selesai": 0, "dibatalkan": 0, "lunas": 0, "belum_dibayar": 0, "dp": 0, "estimated_revenue_le": 0.0}
    for o in orders:
        st = o.get("order_status")
        if st in stats:
            stats[st] += 1
        ps = o.get("payment_status")
        if ps in stats:
            stats[ps] += 1
        if st != "dibatalkan":
            stats["estimated_revenue_le"] += _num(o.get("total_le"))
    return stats

@api_router.get("/admin/analytics/top-configs")
async def admin_top_configs(category: Optional[str] = None, group_by: str = "full",
                            admin: dict = Depends(require_perm("manage_orders"))):
    orders = await db.orders.find({}).to_list(5000)
    agg = {}
    for o in orders:
        if o.get("order_status") == "dibatalkan":
            continue
        for it in o.get("items", [o.get("item")]):
            if not it:
                continue
            if category and it.get("category") != category:
                continue
            cfg = it.get("configuration_snapshot", {}) or {}
            cat = it.get("category")
            if group_by == "length":
                key = f"{cat} — {cfg.get('length', '-')} cm"
            elif group_by == "level":
                key = f"{cat} — {cfg.get('level', '-')} tingkat"
            elif group_by == "type":
                key = f"{cat} — Tipe {cfg.get('type', '-')}"
            elif group_by == "finishing":
                key = f"{cat} — {cfg.get('finishing', '-')}"
            else:
                key = f"{it.get('product_name_snapshot')} — {config_summary(it) or '-'}"
            e = agg.setdefault(key, {"key": key, "category": cat, "config": cfg, "orders": 0, "quantity": 0})
            e["orders"] += 1
            e["quantity"] += int(it.get("quantity", 1))
    rows = sorted(agg.values(), key=lambda x: x["quantity"], reverse=True)
    total_q = sum(r["quantity"] for r in rows) or 1
    for r in rows:
        r["share"] = round(r["quantity"] / total_q * 100, 1)
    return rows[:50]

@api_router.get("/admin/orders/{order_id}")
async def admin_get_order(order_id: str, admin: dict = Depends(require_perm("manage_orders"))):
    order = await db.orders.find_one({"_id": id_query(order_id)})
    if not order:
        raise HTTPException(status_code=404, detail="Pesanan tidak ditemukan")
    order = clean(order)
    order["whatsapp_message"] = build_whatsapp_message(order)
    order["whatsapp_number"] = await get_setting("whatsapp_number", "628XXXXXXXXXX")
    return order

async def record_order_revenue(order):
    if await db.finance_transactions.find_one({"related_order_id": str(order["_id"]), "type": "order_revenue"}):
        return
    now = datetime.now(timezone.utc).isoformat()
    total_le = _num(order.get("total_le"))
    rate = _num(order.get("exchange_rate_idr_per_le"))
    if not rate or rate <= 0:
        rate = _num(await get_setting("exchange_rate_idr_per_le", 357), 357)
    counterpart_idr = _num(order.get("estimated_total_idr"))
    if not counterpart_idr or counterpart_idr <= 0:
        counterpart_idr = round(total_le * rate, 2)
    await db.finance_transactions.insert_one({
        "date": now, "type": "order_revenue", "category": "Penjualan",
        "amount": total_le, "currency": "EGP", "account": "EGP",
        "exchange_rate": rate,
        "counterpart_amount": counterpart_idr,
        "counterpart_currency": "IDR",
        "description": f"Pendapatan otomatis dari pesanan {order.get('order_number')}",
        "related_order_id": str(order["_id"]), "created_by_name": "Sistem",
        "created_at": now, "updated_at": now,
    })

@api_router.patch("/admin/orders/{order_id}")
async def admin_update_order(order_id: str, data: OrderStatusUpdate, admin: dict = Depends(require_perm("manage_orders"))):
    update = audit_fields(admin)
    valid_status = {"pesanan_masuk", "dikonfirmasi", "diproses", "siap", "selesai", "dibatalkan"}
    valid_pay = {"belum_dibayar", "dp", "lunas"}

    existing_order = await db.orders.find_one({"_id": id_query(order_id)})
    if not existing_order:
        raise HTTPException(status_code=404, detail="Pesanan tidak ditemukan")

    if data.order_status is not None:
        if data.order_status not in valid_status:
            raise HTTPException(status_code=400, detail="Status tidak valid")
        update["order_status"] = data.order_status
    if data.payment_status is not None:
        if data.payment_status not in valid_pay:
            raise HTTPException(status_code=400, detail="Status pembayaran tidak valid")
        update["payment_status"] = data.payment_status
    if data.admin_note is not None:
        update["admin_note"] = data.admin_note

    has_price_change = any(x is not None for x in (data.subtotal_le, data.delivery_fee_le, data.discount_le, data.total_le, data.items))
    if has_price_change:
        perms = admin.get("permissions") or {}
        if admin.get("role") != "owner" and not perms.get("access_finance"):
            raise HTTPException(status_code=403, detail="Hanya admin dengan akses keuangan atau owner yang dapat mengubah harga pesanan")

        if data.items is not None:
            update["items"] = data.items
            if len(data.items) > 0:
                update["item"] = data.items[0]
        if data.subtotal_le is not None:
            update["subtotal_le"] = float(data.subtotal_le)
        if data.delivery_fee_le is not None:
            update["delivery_fee_le"] = float(data.delivery_fee_le)
        if data.discount_le is not None:
            update["discount_le"] = float(data.discount_le)

        if data.total_le is not None:
            update["total_le"] = float(data.total_le)
        else:
            cur_sub = update.get("subtotal_le", existing_order.get("subtotal_le", 0))
            cur_del = update.get("delivery_fee_le", existing_order.get("delivery_fee_le", 0))
            cur_disc = update.get("discount_le", existing_order.get("discount_le", 0))
            cur_ref = existing_order.get("referral_discount_le", 0)
            cur_pts = existing_order.get("points_redeemed_le", 0)
            update["total_le"] = max(0.0, float(cur_sub) - float(cur_disc) - float(cur_ref) - float(cur_pts) + float(cur_del))

        rate = existing_order.get("exchange_rate_idr_per_le") or 357
        update["estimated_total_idr"] = round(update["total_le"] * rate)

    await db.orders.update_one({"_id": id_query(order_id)}, {"$set": update})
    order = await db.orders.find_one({"_id": id_query(order_id)})
    effective_pay = data.payment_status if data.payment_status is not None else order.get("payment_status")

    if effective_pay == "lunas":
        now_iso = datetime.now(timezone.utc).isoformat()
        txn = await db.finance_transactions.find_one({"related_order_id": str(order["_id"]), "type": "order_revenue"})
        if txn:
            if has_price_change:
                new_total = _num(order.get("total_le"))
                txn_rate = _num(txn.get("exchange_rate")) or _num(order.get("exchange_rate_idr_per_le"))
                if not txn_rate or txn_rate <= 0:
                    txn_rate = _num(await get_setting("exchange_rate_idr_per_le", 357), 357)
                recomputed_cp = round(new_total * txn_rate, 2)
                await db.finance_transactions.update_one(
                    {"_id": txn["_id"]},
                    {"$set": {
                        "amount": new_total,
                        "exchange_rate": txn_rate,
                        "counterpart_amount": recomputed_cp,
                        "counterpart_currency": "IDR",
                        "updated_at": now_iso,
                        "description": f"Pendapatan otomatis dari pesanan {order.get('order_number')} (disesuaikan)"
                    }}
                )
        else:
            await record_order_revenue(order)
            await award_referral_points(order)
    elif effective_pay in ("belum_dibayar", "dp"):
        await db.finance_transactions.delete_one({"related_order_id": str(order["_id"]), "type": "order_revenue"})
        await reverse_referral_points(order)

    return clean(order)

@api_router.delete("/admin/orders/{order_id}")
async def admin_delete_order(order_id: str, admin: dict = Depends(require_perm("delete_data"))):
    order = await db.orders.find_one({"_id": id_query(order_id)})
    if not order:
        raise HTTPException(status_code=404, detail="Pesanan tidak ditemukan")
    await reverse_referral_points(order)
    await db.orders.delete_one({"_id": id_query(order_id)})
    await db.finance_transactions.delete_many({"related_order_id": order_id, "type": "order_revenue"})
    return {"ok": True}

# --------------------------------------------------------------------------
# Admin: products
# --------------------------------------------------------------------------
def slugify(name):
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return s or str(uuid.uuid4())[:8]

@api_router.post("/admin/products")
async def admin_create_product(payload: Dict[str, Any], admin: dict = Depends(require_perm("modify_products"))):
    now = datetime.now(timezone.utc).isoformat()
    name = payload.get("name", "Produk Baru")
    slug = payload.get("slug") or slugify(name)
    if await db.products.find_one({"slug": slug}):
        slug = f"{slug}-{str(uuid.uuid4())[:4]}"
    doc = {"name": name, "slug": slug, "category": payload.get("category", "custom"),
           "description": payload.get("description", ""), "image_url": payload.get("image_url", ""),
           "active": payload.get("active", True), "configurable": payload.get("configurable", False),
           "starting_price_le": _num(payload.get("starting_price_le", 0)), "pricing": payload.get("pricing", {}),
           "photos": payload.get("photos", []), "representative_photo_id": payload.get("representative_photo_id"),
           "category_cover_image": payload.get("category_cover_image", ""), "cover_mode": payload.get("cover_mode", "manual"),
           "sort_order": payload.get("sort_order", 99), "created_at": now,
           "updated_by_name": admin.get("name"), "updated_at": now}
    result = await db.products.insert_one(doc)
    return clean(await db.products.find_one({"_id": result.inserted_id}))

@api_router.put("/admin/products/{product_id}")
async def admin_update_product(product_id: str, payload: Dict[str, Any], admin: dict = Depends(require_perm("modify_products"))):
    allowed = ["name", "category", "description", "image_url", "active", "configurable",
               "starting_price_le", "pricing", "photos", "representative_photo_id", "sort_order", "slug",
               "category_cover_image", "cover_mode", "photo_weights", "option_notes", "featured_preview_ids"]
    existing = await db.products.find_one({"_id": id_query(product_id)})
    upd = {k: payload[k] for k in allowed if k in payload}
    if "pricing" in upd and not isinstance(upd["pricing"], dict):
        raise HTTPException(status_code=400, detail="Konfigurasi (pricing) harus berupa objek JSON")
    if "option_notes" in upd and isinstance(upd["option_notes"], dict):
        upd["option_notes"] = {k: v for k, v in upd["option_notes"].items() if str(v or "").strip()}
    upd.update(audit_fields(admin))
    if "starting_price_le" in upd:
        upd["starting_price_le"] = _num(upd["starting_price_le"])
    # Record forward-looking price-change history (non-destructive; historical orders keep their own snapshots).
    if existing and ("pricing" in upd or "starting_price_le" in upd):
        old_p = existing.get("pricing"); new_p = upd.get("pricing", old_p)
        old_s = _num(existing.get("starting_price_le")); new_s = _num(upd.get("starting_price_le", old_s))
        if old_p != new_p or old_s != new_s:
            await db.price_history.insert_one({"product_id": product_id, "product_name": existing.get("name"),
                "category": existing.get("category"), "old_pricing": old_p, "new_pricing": new_p,
                "old_starting_price_le": old_s, "new_starting_price_le": new_s,
                "changed_by_name": admin.get("name"), "changed_at": datetime.now(timezone.utc).isoformat()})
    await db.products.update_one({"_id": id_query(product_id)}, {"$set": upd})
    return clean(await db.products.find_one({"_id": id_query(product_id)}))

@api_router.delete("/admin/products/{product_id}")
async def admin_delete_product(product_id: str, admin: dict = Depends(require_perm("delete_data"))):
    await db.products.delete_one({"_id": id_query(product_id)})
    return {"ok": True}

# --------------------------------------------------------------------------
# Uploads
# --------------------------------------------------------------------------
def optimize_image(data: bytes):
    """Resize large images and re-encode to WebP for fast loading. Falls back to original on error."""
    try:
        import io
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        img = img.convert("RGB") if img.mode in ("P", "RGBA", "LA") else img
        max_dim = 1600
        if max(img.size) > max_dim:
            img.thumbnail((max_dim, max_dim), Image.LANCZOS)
        out = io.BytesIO()
        img.save(out, format="WEBP", quality=82, method=4)
        return out.getvalue(), "webp", "image/webp"
    except Exception as e:
        logger.warning(f"Image optimize failed, using original: {e}")
        return None

@api_router.post("/admin/upload")
async def admin_upload(file: UploadFile = File(...), admin: dict = Depends(get_current_admin)):
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "bin").lower()
    if ext not in MIME_TYPES:
        raise HTTPException(status_code=400, detail="Format gambar tidak didukung (jpg, png, webp)")
    data = await file.read()
    if len(data) > 12 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran gambar maksimal 12MB")
    opt = optimize_image(data)
    if opt:
        data, ext, ctype = opt
    else:
        ctype = MIME_TYPES[ext]
    path = f"{APP_NAME}/uploads/{uuid.uuid4()}.{ext}"
    result = put_object(path, data, ctype)
    await db.files.insert_one({"storage_path": result["storage_path"], "original_filename": file.filename,
                               "content_type": ctype, "size": result.get("size", len(data)),
                               "is_deleted": False, "created_at": datetime.now(timezone.utc).isoformat()})
    return {"path": result["storage_path"], "url": f"/api/files/{result['storage_path']}"}

@api_router.get("/files/{path:path}")
async def download_file(path: str):
    record = await db.files.find_one({"storage_path": path, "is_deleted": False})
    if not record:
        raise HTTPException(status_code=404, detail="File tidak ditemukan")
    data, content_type = get_object(path)
    return StarletteResponse(content=data, media_type=record.get("content_type", content_type),
                             headers={"Cache-Control": "public, max-age=86400"})

@api_router.post("/upload-reference")
async def upload_reference(file: UploadFile = File(...)):
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "bin").lower()
    if ext not in MIME_TYPES:
        raise HTTPException(status_code=400, detail="Format gambar tidak didukung (jpg, png, webp)")
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran gambar maksimal 10MB")
    opt = optimize_image(data)
    if opt:
        data, ext, ctype = opt
    else:
        ctype = MIME_TYPES[ext]
    path = f"{APP_NAME}/custom-references/{uuid.uuid4()}.{ext}"
    result = put_object(path, data, ctype)
    await db.files.insert_one({
        "storage_path": result["storage_path"], "original_filename": file.filename,
        "content_type": ctype, "size": result.get("size", len(data)),
        "is_deleted": False, "category": "custom_reference",
        "created_at": datetime.now(timezone.utc).isoformat()
    })
    return {"path": result["storage_path"], "url": f"/api/files/{result['storage_path']}"}

# --------------------------------------------------------------------------
# Custom Requests
# --------------------------------------------------------------------------
@api_router.post("/custom-requests")
async def create_custom_request(data: CustomRequestInput):
    if not data.customer_name or not data.customer_name.strip():
        raise HTTPException(status_code=400, detail="Nama customer wajib diisi")
    phone = normalize_phone(data.customer_phone)
    if not phone.startswith("+"):
        if phone.startswith("08"):
            phone = "+62" + phone[1:]
        elif phone.startswith("62"):
            phone = "+" + phone
        elif phone.startswith("01"):
            phone = "+20" + phone[1:]
        elif phone.startswith("20"):
            phone = "+" + phone
        elif phone:
            phone = "+" + phone
    if not valid_intl_phone(phone):
        raise HTTPException(status_code=400, detail="Nomor telepon WhatsApp tidak valid (format internasional e.g. +62... atau +20...)")
    if not data.furniture_type or not data.furniture_type.strip():
        raise HTTPException(status_code=400, detail="Jenis furniture wajib diisi")

    now = datetime.now(timezone.utc)
    yymm = now.strftime("%y%m")
    count = await db.custom_requests.count_documents({"ticket_number": {"$regex": f"^REQ-{yymm}-"}})
    ticket_number = f"REQ-{yymm}-{count + 1:04d}"

    doc = {
        "ticket_number": ticket_number,
        "customer_name": data.customer_name.strip(),
        "customer_phone": phone,
        "customer_address": (data.customer_address or "").strip(),
        "furniture_type": data.furniture_type.strip(),
        "dimensions": data.dimensions or {},
        "material": (data.material or "").strip(),
        "reference_photos": data.reference_photos or [],
        "budget_estimation_le": float(data.budget_estimation_le) if data.budget_estimation_le is not None else None,
        "notes": (data.notes or "").strip(),
        "status": "baru",
        "admin_notes": "",
        "estimated_price_le": None,
        "converted_order_id": None,
        "created_at": now.isoformat(),
        "updated_at": now.isoformat(),
    }
    result = await db.custom_requests.insert_one(doc)

    wa_num = await get_setting("whatsapp_number", "628XXXXXXXXXX")
    dims_parts = []
    if data.dimensions:
        if data.dimensions.get("length"): dims_parts.append(f"P: {data.dimensions['length']} cm")
        if data.dimensions.get("width"): dims_parts.append(f"L: {data.dimensions['width']} cm")
        if data.dimensions.get("height"): dims_parts.append(f"T: {data.dimensions['height']} cm")
    dims_str = " × ".join(dims_parts) if dims_parts else (data.dimensions.get("notes") if data.dimensions else "")

    lines = [
        "Halo Sogil Furniture, saya ingin mengajukan *Request Custom Furniture*:",
        f"Nomor Tiket: *{ticket_number}*",
        f"Jenis: *{data.furniture_type.strip()}*",
    ]
    if dims_str:
        lines.append(f"Ukuran: {dims_str}")
    if data.material:
        lines.append(f"Pilihan Bahan/Finishing: {data.material.strip()}")
    if data.budget_estimation_le:
        lines.append(f"Estimasi Budget: {fmt_le(data.budget_estimation_le)} LE")
    if data.notes:
        lines.append(f"Catatan: {data.notes.strip()}")
    if data.reference_photos:
        lines.append(f"Foto Referensi: {len(data.reference_photos)} foto terlampir")
    lines.append(f"Nama: {data.customer_name.strip()}")
    lines.append(f"No. WA: {phone}")
    lines.append("\nMohon informasi ketersediaan dan estimasi biayanya. Terima kasih!")

    wa_message = "\n".join(lines)
    clean_num = re.sub(r"[^0-9]", "", wa_num)
    wa_url = f"https://wa.me/{clean_num}?text={urllib.parse.quote(wa_message)}"

    saved = await db.custom_requests.find_one({"_id": result.inserted_id})
    res = clean(saved)
    res["whatsapp_url"] = wa_url
    res["whatsapp_message"] = wa_message
    return res

@api_router.get("/admin/custom-requests")
async def admin_list_custom_requests(
    status: Optional[str] = None,
    q: Optional[str] = None,
    admin: dict = Depends(require_perm("manage_orders"))
):
    query = {}
    if status and status != "all":
        query["status"] = status
    if q:
        regex = {"$regex": re.escape(q), "$options": "i"}
        query["$or"] = [
            {"ticket_number": regex},
            {"customer_name": regex},
            {"customer_phone": regex},
            {"furniture_type": regex},
        ]
    cursor = db.custom_requests.find(query).sort("created_at", -1)
    items = []
    async for doc in cursor:
        items.append(clean(doc))
    return items

@api_router.get("/admin/custom-requests/{req_id}")
async def admin_get_custom_request(req_id: str, admin: dict = Depends(require_perm("manage_orders"))):
    doc = await db.custom_requests.find_one({"_id": id_query(req_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="Request custom tidak ditemukan")
    return clean(doc)

@api_router.patch("/admin/custom-requests/{req_id}")
async def admin_update_custom_request(
    req_id: str,
    data: CustomRequestUpdate,
    admin: dict = Depends(require_perm("manage_orders"))
):
    update = audit_fields(admin)
    if data.status is not None:
        valid_status = {"baru", "diskusi", "disetujui", "dipesan", "ditolak"}
        if data.status not in valid_status:
            raise HTTPException(status_code=400, detail="Status tidak valid")
        update["status"] = data.status
    if data.admin_notes is not None:
        update["admin_notes"] = data.admin_notes
    if data.estimated_price_le is not None:
        update["estimated_price_le"] = float(data.estimated_price_le)

    await db.custom_requests.update_one({"_id": id_query(req_id)}, {"$set": update})
    doc = await db.custom_requests.find_one({"_id": id_query(req_id)})
    return clean(doc)

@api_router.post("/admin/custom-requests/{req_id}/convert-to-order")
async def admin_convert_custom_to_order(
    req_id: str,
    data: ConvertCustomToOrderInput,
    admin: dict = Depends(require_perm("manage_orders"))
):
    custom_req = await db.custom_requests.find_one({"_id": id_query(req_id)})
    if not custom_req:
        raise HTTPException(status_code=404, detail="Request custom tidak ditemukan")

    order_num = await generate_order_number()
    rate = _num(await get_setting("exchange_rate_idr_per_le", 357), 357)
    now = datetime.now(timezone.utc).isoformat()
    subtotal = float(data.subtotal_le)
    delivery_fee = float(data.delivery_fee_le or 0.0)
    total_le = max(0.0, subtotal + delivery_fee)

    zone_name = None
    if data.delivery_zone_id:
        z = await db.delivery_zones.find_one({"_id": id_query(data.delivery_zone_id)})
        if z:
            zone_name = z.get("name")

    item_snapshot = {
        "product_id": str(custom_req["_id"]),
        "product_name_snapshot": f"Custom: {custom_req.get('furniture_type', 'Furniture')}",
        "category": "custom",
        "configuration_snapshot": {
            "custom_size": True,
            "ticket_number": custom_req.get("ticket_number"),
            "furniture_type": custom_req.get("furniture_type"),
            "dimensions": custom_req.get("dimensions", {}),
            "material": custom_req.get("material", ""),
            "reference_photos": custom_req.get("reference_photos", []),
            "custom_note": custom_req.get("notes", ""),
        },
        "quantity": 1,
        "unit_price_le": subtotal,
        "subtotal_le": subtotal,
    }

    order_doc = {
        "order_number": order_num,
        "customer_name": custom_req.get("customer_name"),
        "customer_phone": custom_req.get("customer_phone"),
        "customer_address": data.customer_address or custom_req.get("customer_address") or "Dikonfirmasi via WA",
        "customer_maps_url": data.customer_maps_url or "",
        "delivery_method": data.delivery_method,
        "delivery_zone_id": data.delivery_zone_id,
        "delivery_zone_name": zone_name,
        "delivery_fee_le": delivery_fee,
        "payment_method": data.payment_method,
        "payment_status": "belum_dibayar",
        "order_status": "dikonfirmasi",
        "notes": f"[Tiket {custom_req.get('ticket_number')}] {data.notes or custom_req.get('notes', '')}",
        "admin_note": f"Dikonversi dari tiket {custom_req.get('ticket_number')} oleh {admin.get('name')}",
        "items": [item_snapshot],
        "item": item_snapshot,
        "subtotal_le": subtotal,
        "discount_code": None,
        "discount_percentage": 0,
        "discount_le": 0,
        "referral": None,
        "referral_discount_le": 0,
        "points_redeemed_le": 0,
        "customer_id": None,
        "customer_username": None,
        "total_le": total_le,
        "exchange_rate_idr_per_le": rate,
        "estimated_total_idr": round(total_le * rate),
        "rate_timestamp": now,
        "requires_admin_confirmation": False,
        "custom_request_id": str(custom_req["_id"]),
        "created_at": now,
        "updated_at": now,
    }

    order_result = await db.orders.insert_one(order_doc)
    order_id = str(order_result.inserted_id)

    await db.custom_requests.update_one(
        {"_id": custom_req["_id"]},
        {"$set": {
            "status": "dipesan",
            "converted_order_id": order_id,
            "estimated_price_le": subtotal,
            "updated_at": now,
        }}
    )

    created_order = await db.orders.find_one({"_id": order_result.inserted_id})
    return clean(created_order)

@api_router.delete("/admin/custom-requests/{req_id}")
async def admin_delete_custom_request(req_id: str, admin: dict = Depends(require_perm("delete_data"))):
    res = await db.custom_requests.delete_one({"_id": id_query(req_id)})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Request custom tidak ditemukan")
    return {"ok": True}

@api_router.post("/admin/custom-collection/publish")
async def admin_publish_to_custom_collection(data: PublishCustomCollectionInput, admin: dict = Depends(require_perm("modify_products"))):
    title = data.title.strip()
    if not title:
        raise HTTPException(status_code=400, detail="Nama desain harus diisi")

    custom_prod = await db.products.find_one({"category": "custom"})
    if not custom_prod:
        custom_prod = await db.products.find_one({"slug": {"$in": ["koleksi-custom", "pesanan-custom"]}})
    if not custom_prod:
        raise HTTPException(status_code=404, detail="Produk Koleksi Custom tidak ditemukan di database")

    prod_id = custom_prod["_id"]
    pricing = custom_prod.get("pricing") or {}
    groups = pricing.get("groups") or []
    photos = list(custom_prod.get("photos") or [])
    design_details = dict(custom_prod.get("design_details") or {})

    desain_group = None
    for g in groups:
        if g.get("key") == "desain":
            desain_group = g
            break
    if not desain_group:
        desain_group = {"key": "desain", "label": "Pilih Desain Custom", "options": []}
        groups.append(desain_group)

    if title not in desain_group.get("options", []):
        desain_group.setdefault("options", []).append(title)

    for url in data.photo_urls:
        url_clean = url.strip()
        if not url_clean:
            continue
        exists = any(ph.get("url") == url_clean and ph.get("attributes", {}).get("desain") == title for ph in photos)
        if not exists:
            photos.append({
                "url": url_clean,
                "file_id": None,
                "attributes": {"desain": title}
            })

    design_details[title] = {
        "price": float(data.price_le),
        "spesifikasi": data.spesifikasi or "",
        "published_at": datetime.now(timezone.utc).isoformat(),
        "published_by": admin.get("name") or admin.get("email")
    }

    pricing["groups"] = groups

    await db.products.update_one(
        {"_id": prod_id},
        {"$set": {
            "pricing": pricing,
            "photos": photos,
            "design_details": design_details,
            "updated_at": datetime.now(timezone.utc).isoformat()
        }}
    )

    if data.order_id:
        await db.orders.update_one(
            {"_id": id_query(data.order_id)},
            {"$set": {
                "custom_collection_published": True,
                "custom_collection_title": title,
                "custom_collection_published_at": datetime.now(timezone.utc).isoformat()
            }}
        )

    if data.request_id:
        await db.custom_requests.update_one(
            {"_id": id_query(data.request_id)},
            {"$set": {
                "custom_collection_published": True,
                "custom_collection_title": title,
                "custom_collection_published_at": datetime.now(timezone.utc).isoformat()
            }}
        )

    return {"ok": True, "title": title, "photos_count": len(photos)}

# --------------------------------------------------------------------------
# Admin: delivery zones
# --------------------------------------------------------------------------
@api_router.post("/admin/delivery-zones")
async def admin_create_zone(data: DeliveryZoneInput, admin: dict = Depends(require_perm("manage_settings"))):
    doc = {"name": data.name, "fee_le": _num(data.fee_le), "active": data.active, "created_at": datetime.now(timezone.utc).isoformat()}
    result = await db.delivery_zones.insert_one(doc)
    return clean(await db.delivery_zones.find_one({"_id": result.inserted_id}))

@api_router.put("/admin/delivery-zones/{zone_id}")
async def admin_update_zone(zone_id: str, data: DeliveryZoneInput, admin: dict = Depends(require_perm("manage_settings"))):
    await db.delivery_zones.update_one({"_id": id_query(zone_id)}, {"$set": {"name": data.name, "fee_le": _num(data.fee_le), "active": data.active}})
    return clean(await db.delivery_zones.find_one({"_id": id_query(zone_id)}))

@api_router.delete("/admin/delivery-zones/{zone_id}")
async def admin_delete_zone(zone_id: str, admin: dict = Depends(require_perm("manage_settings"))):
    await db.delivery_zones.delete_one({"_id": id_query(zone_id)})
    return {"ok": True}

# --------------------------------------------------------------------------
# Admin: settings
# --------------------------------------------------------------------------
@api_router.get("/admin/settings")
async def admin_get_settings(admin: dict = Depends(require_perm("manage_settings"))):
    info = await store_info()
    meta = {}
    for k in ["store_name", "whatsapp_number", "exchange_rate_idr_per_le", "logo_url"]:
        doc = await db.settings.find_one({"key": k})
        if doc and doc.get("updated_by_name"):
            meta[k] = {"updated_by_name": doc.get("updated_by_name"), "updated_at": doc.get("updated_at")}
    info["_audit"] = meta
    return info

@api_router.put("/admin/settings")
async def admin_update_settings(payload: Dict[str, Any], admin: dict = Depends(require_perm("manage_settings"))):
    allowed = ["store_name", "app_name", "tagline", "store_address", "store_maps_url",
               "whatsapp_number", "exchange_rate_idr_per_le", "bank_info", "logo_url",
               "instagram", "facebook", "tiktok", "email", "referral_program_enabled", "point_redeem_max_pct"]
    for key in allowed:
        if key in payload:
            val = payload[key]
            if key == "exchange_rate_idr_per_le":
                val = _num(val, 357)
            await set_setting(key, val, admin)
    return await store_info()

# --------------------------------------------------------------------------
# Admin: manage admins (owner / manage_admins)
# --------------------------------------------------------------------------
@api_router.get("/admin/admins")
async def list_admins(admin: dict = Depends(require_perm("manage_admins"))):
    admins = await db.admins.find({}).to_list(200)
    out = []
    for a in admins:
        a = clean(a)
        if a.get("role") == "owner":
            a["permissions"] = default_permissions("owner")
        else:
            a.setdefault("permissions", default_permissions(a.get("role", "admin")))
        out.append(a)
    role_priority = {"owner": 0, "manager": 1, "admin": 2, "employee": 3}
    out.sort(key=lambda x: (role_priority.get(x.get("role", "employee"), 99), (x.get("name") or "").lower()))
    return out

@api_router.post("/admin/admins")
async def create_admin(data: AdminCreateInput, admin: dict = Depends(require_owner())):
    email = data.email.strip().lower()
    if data.role not in ACCOUNT_TYPES:
        raise HTTPException(status_code=400, detail="Tipe akun tidak valid")
    if data.role == "owner":
        raise HTTPException(status_code=400, detail="Tidak dapat membuat akun Owner tambahan")
    if await db.admins.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah digunakan")
    if len(data.password) < 6:
        raise HTTPException(status_code=400, detail="Kata sandi minimal 6 karakter")
    perms = default_permissions(data.role)
    if data.permissions:
        for k in PERMISSION_KEYS:
            if k in data.permissions:
                perms[k] = bool(data.permissions[k])
    perms["manage_admins"] = False  # only owner
    doc = {"name": data.name.strip(), "email": email, "password_hash": hash_password(data.password),
           "role": data.role, "permissions": perms, "created_at": datetime.now(timezone.utc).isoformat(),
           "created_by_name": admin.get("name")}
    result = await db.admins.insert_one(doc)
    return clean(await db.admins.find_one({"_id": result.inserted_id}))

@api_router.put("/admin/admins/{admin_id}")
async def update_admin(admin_id: str, data: AdminUpdateInput, admin: dict = Depends(require_owner())):
    target = await db.admins.find_one({"_id": id_query(admin_id)})
    if not target:
        raise HTTPException(status_code=404, detail="Akun tidak ditemukan")
    if target.get("role") == "owner":
        raise HTTPException(status_code=400, detail="Akun Owner tidak dapat diubah dari sini")
    upd = {}
    if data.name is not None:
        upd["name"] = data.name.strip()
    if data.status in ("active", "inactive"):
        upd["status"] = data.status
    if data.email is not None:
        new_email = data.email.strip().lower()
        if await db.admins.find_one({"email": new_email, "_id": {"$ne": oid(admin_id)}}):
            raise HTTPException(status_code=400, detail="Email sudah digunakan")
        upd["email"] = new_email
    if data.password:
        if len(data.password) < 6:
            raise HTTPException(status_code=400, detail="Kata sandi minimal 6 karakter")
        upd["password_hash"] = hash_password(data.password)
    if data.role is not None:
        if data.role not in ACCOUNT_TYPES or data.role == "owner":
            raise HTTPException(status_code=400, detail="Tipe akun tidak valid")
        upd["role"] = data.role
    if data.permissions is not None:
        perms = target.get("permissions", default_permissions(target.get("role", "admin")))
        for k in PERMISSION_KEYS:
            if k in data.permissions:
                perms[k] = bool(data.permissions[k])
        perms["manage_admins"] = False
        upd["permissions"] = perms
    await db.admins.update_one({"_id": id_query(admin_id)}, {"$set": upd})
    return clean(await db.admins.find_one({"_id": id_query(admin_id)}))

@api_router.delete("/admin/admins/{admin_id}")
async def delete_admin(admin_id: str, admin: dict = Depends(require_owner())):
    target = await db.admins.find_one({"_id": id_query(admin_id)})
    if not target:
        raise HTTPException(status_code=404, detail="Akun tidak ditemukan")
    if target.get("role") == "owner":
        raise HTTPException(status_code=400, detail="Akun Owner tidak dapat dihapus")
    await db.admins.delete_one({"_id": id_query(admin_id)})
    return {"ok": True}

# --------------------------------------------------------------------------
# Finance
# --------------------------------------------------------------------------
DEFAULT_FINANCE_CATEGORIES = {
    "income": ["Penjualan", "Pendapatan Lain"],
    "expense": ["Material", "Pekerja", "Biaya Sewa Ruang", "Naql", "Parts", "Pemotongan", "Transport",
                "Peralatan & Perkakas", "Perbaikan & Pemeliharaan", "Konsumsi Operasional", "Air & Listrik",
                "Mesin & Kendaraan", "Pengembangan Aplikasi Web", "Marketing"],
}

@api_router.get("/admin/finance/categories")
async def finance_categories(admin: dict = Depends(require_perm("access_finance"))):
    custom = await db.finance_categories.find({"status": "active"}).to_list(200)
    out = {"income": list(DEFAULT_FINANCE_CATEGORIES["income"]), "expense": list(DEFAULT_FINANCE_CATEGORIES["expense"])}
    for c in custom:
        typ = c.get("type", "expense")
        if typ in out and c["name"] not in out[typ]:
            out[typ].append(c["name"])
    return out

async def _compute_balances():
    bal = {"IDR": 0.0, "EGP": 0.0}
    async for t in db.finance_transactions.find({}):
        typ = t.get("type"); cur = t.get("currency", "EGP"); amt = _num(t.get("amount"))
        if typ in ("income", "order_revenue"):
            bal[t.get("account", cur)] = bal.get(t.get("account", cur), 0) + amt
        elif typ == "expense":
            bal[t.get("account", cur)] = bal.get(t.get("account", cur), 0) - amt
        elif typ == "transfer":
            bal[t.get("from_account", "IDR")] = bal.get(t.get("from_account", "IDR"), 0) - _num(t.get("from_amount"))
            bal[t.get("to_account", "EGP")] = bal.get(t.get("to_account", "EGP"), 0) + _num(t.get("to_amount"))
        elif typ == "balance_adjustment":
            bal[t.get("account", cur)] = bal.get(t.get("account", cur), 0) + _num(t.get("amount"))
    return bal

@api_router.get("/admin/finance/accounts")
async def finance_accounts(admin: dict = Depends(require_perm("access_finance"))):
    return {"balances": await _compute_balances()}

@api_router.get("/admin/finance/transactions")
async def finance_list(currency: Optional[str] = None, type: Optional[str] = None,
                       start: Optional[str] = None, end: Optional[str] = None, q: Optional[str] = None,
                       admin: dict = Depends(require_perm("access_finance"))):
    conds = []
    if currency:
        conds.append({"$or": [{"currency": currency}, {"from_account": currency}, {"to_account": currency}]})
    if type:
        conds.append({"type": type})
    if start or end:
        d = {}
        if start:
            d["$gte"] = start
        if end:
            d["$lte"] = end + "T23:59:59"
        conds.append({"date": d})
    if q and q.strip():
        qs = q.strip()
        rx = {"$regex": re.escape(qs), "$options": "i"}
        ors = [{"description": rx}, {"category": rx}, {"type": rx}, {"classification": rx},
               {"created_by_name": rx}, {"recipient_employee_name": rx}]
        try:
            ors.append({"amount": float(qs)})
        except ValueError:
            pass
        conds.append({"$or": ors})
    query = {"$and": conds} if conds else {}
    txns = await db.finance_transactions.find(query).sort("date", -1).to_list(2000)
    return [clean(t) for t in txns]

async def category_classification(name, typ):
    """Revenue/Cost vs Transfer classification for a category (snapshot onto each txn)."""
    c = await db.finance_categories.find_one({"name": name, "type": typ})
    if c and c.get("classification") in ("revenue", "cost", "transfer"):
        return c["classification"]
    return "revenue" if typ == "income" else "cost"

def _is_wage_category(name):
    n = (name or "").lower()
    return "upah" in n or "wage" in n or "pekerja" in n

@api_router.post("/admin/finance/transactions")
async def finance_create(data: FinanceTxnInput, admin: dict = Depends(require_perm("access_finance"))):
    if data.type not in ("income", "expense"):
        raise HTTPException(status_code=400, detail="Tipe transaksi tidak valid")
    if data.currency not in ("IDR", "EGP"):
        raise HTTPException(status_code=400, detail="Mata uang tidak valid")
    now = datetime.now(timezone.utc).isoformat()
    classification = await category_classification(data.category, data.type)

    settings_rate = _num(await get_setting("exchange_rate_idr_per_le", 357), 357)
    if data.exchange_rate is not None and _num(data.exchange_rate) > 0:
        rate = _num(data.exchange_rate)
    else:
        rate = settings_rate

    primary_amt = _num(data.amount)
    if data.currency == "EGP":
        counterpart_cur = "IDR"
        if data.counterpart_amount is not None and _num(data.counterpart_amount) > 0 and primary_amt == 0:
            primary_amt = round(_num(data.counterpart_amount) / rate, 2) if rate > 0 else 0.0
            counterpart_amt = round(_num(data.counterpart_amount), 2)
        else:
            counterpart_amt = round(primary_amt * rate, 2)
    else:
        counterpart_cur = "EGP"
        if data.counterpart_amount is not None and _num(data.counterpart_amount) > 0 and primary_amt == 0:
            primary_amt = round(_num(data.counterpart_amount) * rate, 2)
            counterpart_amt = round(_num(data.counterpart_amount), 2)
        else:
            counterpart_amt = round(primary_amt / rate, 2) if rate > 0 else 0.0

    doc = {"date": data.date or now, "type": data.type, "category": data.category,
           "amount": primary_amt, "currency": data.currency, "account": data.currency,
           "exchange_rate": rate,
           "counterpart_amount": counterpart_amt,
           "counterpart_currency": counterpart_cur,
           "classification": classification,
           "description": data.description or "", "related_order_id": None,
           "created_by_id": admin.get("id"), "created_by_name": admin.get("name"),
           "updated_by_name": admin.get("name"), "created_at": now, "updated_at": now}
    is_wage = data.type == "expense" and _is_wage_category(data.category)
    if is_wage and data.recipient_employee_id:
        emp = await db.admins.find_one({"_id": id_query(data.recipient_employee_id)})
        if not emp:
            raise HTTPException(status_code=400, detail="Karyawan penerima tidak ditemukan")
        doc["recipient_employee_id"] = data.recipient_employee_id
        doc["recipient_employee_name"] = emp.get("name")
    result = await db.finance_transactions.insert_one(doc)
    tid = str(result.inserted_id)
    if doc.get("recipient_employee_id"):
        await db.employee_wages.insert_one({"employee_id": doc["recipient_employee_id"], "employee_name": doc["recipient_employee_name"],
            "amount": primary_amt, "currency": data.currency, "date": doc["date"], "category": data.category,
            "description": data.description or "", "transaction_id": tid, "recorded_by_name": admin.get("name"), "created_at": now})
    return clean(await db.finance_transactions.find_one({"_id": result.inserted_id}))

@api_router.put("/admin/finance/transactions/{txn_id}")
async def finance_update(txn_id: str, data: FinanceTxnInput, admin: dict = Depends(require_perm("access_finance"))):
    t = await db.finance_transactions.find_one({"_id": id_query(txn_id)})
    if not t:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")
    if t.get("type") == "order_revenue":
        raise HTTPException(status_code=400, detail="Pendapatan otomatis pesanan tidak dapat diubah manual")

    existing_rate = _num(t.get("exchange_rate"))
    settings_rate = _num(await get_setting("exchange_rate_idr_per_le", 357), 357)
    base_rate = existing_rate if (existing_rate and existing_rate > 0) else settings_rate

    if data.exchange_rate is not None and _num(data.exchange_rate) > 0:
        rate = _num(data.exchange_rate)
    else:
        rate = base_rate

    existing_amt = _num(t.get("amount"))
    existing_cp = _num(t.get("counterpart_amount"))
    input_amt = _num(data.amount)
    input_cp = _num(data.counterpart_amount) if data.counterpart_amount is not None else None

    if input_cp is not None and input_amt == existing_amt and input_cp != existing_cp:
        if data.currency == "EGP":
            primary_amt = round(input_cp / rate, 2) if rate > 0 else 0.0
            counterpart_amt = input_cp
            counterpart_cur = "IDR"
        else:
            primary_amt = round(input_cp * rate, 2)
            counterpart_amt = input_cp
            counterpart_cur = "EGP"
    else:
        primary_amt = input_amt
        if data.currency == "EGP":
            counterpart_cur = "IDR"
            counterpart_amt = round(primary_amt * rate, 2)
        else:
            counterpart_cur = "EGP"
            counterpart_amt = round(primary_amt / rate, 2) if rate > 0 else 0.0

    upd = {"date": data.date or t.get("date"), "type": data.type, "category": data.category,
           "amount": primary_amt, "currency": data.currency, "account": data.currency,
           "exchange_rate": rate,
           "counterpart_amount": counterpart_amt,
           "counterpart_currency": counterpart_cur,
           "classification": await category_classification(data.category, data.type),
           "description": data.description or "", "updated_by_name": admin.get("name"),
           "updated_at": datetime.now(timezone.utc).isoformat()}
    # Re-sync wage linkage: recompute wage-ness/recipient so employee history stays consistent on edit.
    is_wage = data.type == "expense" and _is_wage_category(data.category)
    upd["recipient_employee_id"] = None
    upd["recipient_employee_name"] = None
    if is_wage and data.recipient_employee_id:
        emp = await db.admins.find_one({"_id": id_query(data.recipient_employee_id)})
        if not emp:
            raise HTTPException(status_code=400, detail="Karyawan penerima tidak ditemukan")
        upd["recipient_employee_id"] = data.recipient_employee_id
        upd["recipient_employee_name"] = emp.get("name")
    await db.finance_transactions.update_one({"_id": id_query(txn_id)}, {"$set": upd})
    await db.employee_wages.delete_many({"transaction_id": txn_id})
    if upd["recipient_employee_id"]:
        await db.employee_wages.insert_one({"employee_id": upd["recipient_employee_id"], "employee_name": upd["recipient_employee_name"],
            "amount": primary_amt, "currency": data.currency, "date": upd["date"], "category": data.category,
            "description": data.description or "", "transaction_id": txn_id, "recorded_by_name": admin.get("name"),
            "created_at": datetime.now(timezone.utc).isoformat()})
    return clean(await db.finance_transactions.find_one({"_id": id_query(txn_id)}))

@api_router.delete("/admin/finance/transactions/{txn_id}")
async def finance_delete(txn_id: str, admin: dict = Depends(require_perm("access_finance"))):
    await db.employee_wages.delete_many({"transaction_id": txn_id})
    await db.finance_transactions.delete_one({"_id": id_query(txn_id)})
    return {"ok": True}

@api_router.post("/admin/finance/transfer")
async def finance_transfer(data: TransferInput, admin: dict = Depends(require_perm("access_finance"))):
    if data.from_account not in ("IDR", "EGP") or data.to_account not in ("IDR", "EGP") or data.from_account == data.to_account:
        raise HTTPException(status_code=400, detail="Akun transfer tidak valid")
    if _num(data.from_amount) <= 0 or _num(data.to_amount) <= 0:
        raise HTTPException(status_code=400, detail="Jumlah transfer harus lebih dari 0")
    now = datetime.now(timezone.utc).isoformat()
    doc = {"date": data.date or now, "type": "transfer", "category": "Transfer/Konversi",
           "from_account": data.from_account, "to_account": data.to_account,
           "from_amount": _num(data.from_amount), "to_amount": _num(data.to_amount),
           "exchange_rate": _num(data.exchange_rate) if data.exchange_rate else None,
           "description": data.description or "", "created_by_name": admin.get("name"),
           "updated_by_name": admin.get("name"), "created_at": now, "updated_at": now}
    result = await db.finance_transactions.insert_one(doc)
    return clean(await db.finance_transactions.find_one({"_id": result.inserted_id}))

def _period_range(period, start, end):
    now = datetime.now(timezone.utc)
    if period == "custom" and start and end:
        return start, end + "T23:59:59", None, None
    if period == "this_month":
        s = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        ps = (s - timedelta(days=1)).replace(day=1)
        return s.isoformat(), now.isoformat(), ps.isoformat(), s.isoformat()
    if period == "last_month":
        first = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        last_end = first - timedelta(seconds=1)
        s = last_end.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        pend = s - timedelta(seconds=1); ps = pend.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        return s.isoformat(), last_end.isoformat(), ps.isoformat(), s.isoformat()
    if period == "this_year":
        s = now.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
        ps = s.replace(year=s.year - 1)
        return s.isoformat(), now.isoformat(), ps.isoformat(), s.isoformat()
    if period == "last_year":
        s = now.replace(year=now.year - 1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
        e = now.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0) - timedelta(seconds=1)
        ps = s.replace(year=s.year - 1)
        return s.isoformat(), e.isoformat(), ps.isoformat(), s.isoformat()
    if period == "last_3_months":
        m = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        y, mo = m.year, m.month - 2
        while mo <= 0:
            mo += 12; y -= 1
        s = m.replace(year=y, month=mo)
        return s.isoformat(), now.isoformat(), None, None
    return None, None, None, None

async def _stats_for(s, e):
    q = {}
    if s or e:
        q["date"] = {}
        if s:
            q["date"]["$gte"] = s
        if e:
            q["date"]["$lte"] = e
    res = {"IDR": {"revenue": 0.0, "cost": 0.0}, "EGP": {"revenue": 0.0, "cost": 0.0}}
    async for t in db.finance_transactions.find(q):
        cur = t.get("currency", "EGP"); typ = t.get("type"); amt = _num(t.get("amount")); cls = t.get("classification")
        if cur not in res:
            continue
        if typ == "order_revenue":
            res[cur]["revenue"] += amt
        elif typ == "income" and cls in (None, "revenue"):
            res[cur]["revenue"] += amt
        elif typ == "expense" and cls in (None, "cost"):
            res[cur]["cost"] += amt
    for cur in res:
        r = res[cur]
        r["operating_profit"] = r["revenue"] - r["cost"]
        r["cash_flow"] = r["revenue"] - r["cost"]
    return res

@api_router.get("/admin/finance/stats")
async def finance_stats(period: str = "this_month", start: Optional[str] = None, end: Optional[str] = None,
                        admin: dict = Depends(require_perm("access_finance"))):
    s, e, ps, pe = _period_range(period, start, end)
    current = await _stats_for(s, e)
    balances = await _compute_balances()
    comparison = None
    if ps and pe:
        prev = await _stats_for(ps, pe)
        comparison = {}
        for cur in ("IDR", "EGP"):
            comparison[cur] = {}
            for metric in ("revenue", "cost", "operating_profit"):
                cv = current[cur][metric]; pv = prev[cur][metric]
                pct = ((cv - pv) / abs(pv) * 100) if pv else (100.0 if cv else 0.0)
                comparison[cur][metric] = round(pct, 1)
    return {"period": period, "range": {"start": s, "end": e}, "current": current,
            "comparison": comparison, "balances": balances}

async def validate_discount(code, subtotal):
    if not code:
        return 0.0, None, None
    d = await db.discounts.find_one({"code": code.strip().upper()})
    if not d:
        return 0.0, None, "Kode diskon tidak ditemukan"
    now = datetime.now(timezone.utc).isoformat()
    if d.get("status") != "active":
        return 0.0, None, "Kode diskon tidak aktif"
    if d.get("start_date") and now < d["start_date"]:
        return 0.0, None, "Kode diskon belum berlaku"
    if d.get("end_date") and now > (d["end_date"] + "T23:59:59"):
        return 0.0, None, "Kode diskon sudah kedaluwarsa"
    if d.get("max_claims") and _num(d.get("claims", 0)) >= _num(d["max_claims"]):
        return 0.0, None, "Kuota kode diskon habis"
    disc = subtotal * _num(d.get("percentage")) / 100.0
    mx = _num(d.get("max_amount"))
    if mx > 0:
        disc = min(disc, mx)
    return round(disc, 2), d, None

@api_router.post("/discounts/validate")
async def discount_validate(payload: Dict[str, Any]):
    disc, d, err = await validate_discount(payload.get("code"), _num(payload.get("subtotal")))
    if err:
        return {"valid": False, "message": err, "discount_amount": 0}
    return {"valid": True, "discount_amount": disc, "percentage": _num(d.get("percentage")), "code": d.get("code")}

@api_router.get("/admin/discounts")
async def list_discounts(admin: dict = Depends(require_perm("manage_settings"))):
    return [clean(d) for d in await db.discounts.find({}).sort("created_at", -1).to_list(500)]

@api_router.post("/admin/discounts")
async def create_discount(payload: Dict[str, Any], admin: dict = Depends(require_perm("manage_settings"))):
    code = (payload.get("code") or "").strip().upper()
    if not code:
        raise HTTPException(status_code=400, detail="Kode diskon wajib diisi")
    if await db.discounts.find_one({"code": code}):
        raise HTTPException(status_code=400, detail="Kode diskon sudah ada")
    now = datetime.now(timezone.utc).isoformat()
    doc = {"code": code, "name": payload.get("name", ""), "percentage": max(0.0, min(100.0, _num(payload.get("percentage")))),
           "max_amount": _num(payload.get("max_amount")), "max_claims": int(_num(payload.get("max_claims"))),
           "claims": 0, "start_date": payload.get("start_date") or None, "end_date": payload.get("end_date") or None,
           "status": payload.get("status", "active"), "created_at": now, "updated_by_name": admin.get("name")}
    r = await db.discounts.insert_one(doc)
    return clean(await db.discounts.find_one({"_id": r.inserted_id}))

@api_router.put("/admin/discounts/{did}")
async def update_discount(did: str, payload: Dict[str, Any], admin: dict = Depends(require_perm("manage_settings"))):
    upd = {}
    for k in ["name", "percentage", "max_amount", "max_claims", "start_date", "end_date", "status"]:
        if k in payload:
            upd[k] = payload[k]
    for k in ["percentage", "max_amount", "max_claims"]:
        if k in upd:
            upd[k] = _num(upd[k])
    if "percentage" in upd:
        upd["percentage"] = max(0.0, min(100.0, upd["percentage"]))
    upd["updated_by_name"] = admin.get("name")
    await db.discounts.update_one({"_id": id_query(did)}, {"$set": upd})
    return clean(await db.discounts.find_one({"_id": id_query(did)}))

@api_router.delete("/admin/discounts/{did}")
async def delete_discount(did: str, admin: dict = Depends(require_perm("manage_settings"))):
    await db.discounts.delete_one({"_id": id_query(did)})
    return {"ok": True}

@api_router.get("/admin/finance/custom-categories")
async def list_custom_categories(admin: dict = Depends(require_perm("access_finance"))):
    return [clean(c) for c in await db.finance_categories.find({}).to_list(200)]

@api_router.post("/admin/finance/custom-categories")
async def create_custom_category(payload: Dict[str, Any], admin: dict = Depends(require_perm("access_finance"))):
    name = (payload.get("name") or "").strip()
    typ = payload.get("type", "expense")
    if not name or typ not in ("income", "expense"):
        raise HTTPException(status_code=400, detail="Nama & tipe kategori wajib benar")
    valid = {"income": ["revenue", "transfer"], "expense": ["cost", "transfer"]}
    classification = payload.get("classification")
    if classification not in valid[typ]:
        classification = "revenue" if typ == "income" else "cost"
    doc = {"name": name, "type": typ, "classification": classification, "status": payload.get("status", "active"),
           "created_at": datetime.now(timezone.utc).isoformat(), "created_by_name": admin.get("name")}
    r = await db.finance_categories.insert_one(doc)
    return clean(await db.finance_categories.find_one({"_id": r.inserted_id}))

@api_router.put("/admin/finance/custom-categories/{cid}")
async def update_custom_category(cid: str, payload: Dict[str, Any], admin: dict = Depends(require_perm("access_finance"))):
    upd = {}
    if payload.get("classification") in ("revenue", "cost", "transfer"):
        upd["classification"] = payload["classification"]
    if payload.get("status") in ("active", "inactive"):
        upd["status"] = payload["status"]
    if upd:
        await db.finance_categories.update_one({"_id": id_query(cid)}, {"$set": upd})
    return clean(await db.finance_categories.find_one({"_id": id_query(cid)}))

@api_router.delete("/admin/finance/custom-categories/{cid}")
async def delete_custom_category(cid: str, admin: dict = Depends(require_perm("access_finance"))):
    # Soft-delete: historical transactions keep their category/classification; just stop offering it for new txns.
    await db.finance_categories.update_one({"_id": id_query(cid)}, {"$set": {"status": "inactive"}})
    return {"ok": True}

@api_router.get("/admin/employees")
async def list_employees(admin: dict = Depends(require_perm("access_finance"))):
    return [clean(a) for a in await db.admins.find({"status": {"$ne": "inactive"}}).sort("name", 1).to_list(200)]

@api_router.get("/admin/employees/wages")
async def employee_wages_summary(admin: dict = Depends(require_perm("access_finance"))):
    now = datetime.now(timezone.utc)
    mkey = now.strftime("%Y-%m"); ykey = now.strftime("%Y")
    out = []
    for e in await db.admins.find({}).sort("name", 1).to_list(200):
        eid = str(e["_id"])
        # Source of truth: finance_transactions
        txns = await db.finance_transactions.find({
            "type": "expense",
            "recipient_employee_id": eid
        }).sort("date", -1).to_list(1000)

        # Fallback to employee_wages mirror if no transactions found (for legacy compatibility)
        if not txns:
            txns = await db.employee_wages.find({"employee_id": eid}).sort("date", -1).to_list(500)

        totals, month, year = {}, {}, {}
        history = []
        for w in txns:
            if not _is_wage_category(w.get("category", "")):
                continue
            c = w.get("currency", "EGP"); amt = _num(w.get("amount")); d = (w.get("date") or "")
            totals[c] = totals.get(c, 0) + amt
            if d[:7] == mkey:
                month[c] = month.get(c, 0) + amt
            if d[:4] == ykey:
                year[c] = year.get(c, 0) + amt
            history.append({
                "id": str(w["_id"]),
                "amount": amt,
                "currency": c,
                "date": d,
                "category": w.get("category", "Pekerja"),
                "description": w.get("description", ""),
                "recorded_by_name": w.get("recorded_by_name") or w.get("created_by_name") or "-",
                "created_at": w.get("created_at")
            })
        out.append({"id": eid, "name": e.get("name"), "email": e.get("email"), "role": e.get("role"),
                    "status": e.get("status", "active"), "totals": totals, "month": month, "year": year,
                    "count": len(history), "history": history})
    return out

@api_router.post("/admin/finance/balance-adjust")
async def balance_adjust(payload: Dict[str, Any], admin: dict = Depends(require_perm("access_finance"))):
    account = payload.get("account")
    if account not in ("IDR", "EGP"):
        raise HTTPException(status_code=400, detail="Akun tidak valid")
    balances = await _compute_balances()
    prev = balances.get(account, 0.0)
    new_balance = _num(payload.get("new_balance"))
    delta = new_balance - prev
    now = datetime.now(timezone.utc).isoformat()

    rate = _num(payload.get("exchange_rate"))
    if not rate or rate <= 0:
        rate = _num(await get_setting("exchange_rate_idr_per_le", 357), 357)

    if account == "EGP":
        counterpart_cur = "IDR"
        counterpart_amt = round(delta * rate, 2)
    else:
        counterpart_cur = "EGP"
        counterpart_amt = round(delta / rate, 2) if rate > 0 else 0.0

    doc = {"date": payload.get("date") or now, "type": "balance_adjustment", "category": "Penyesuaian Saldo",
           "amount": delta, "currency": account, "account": account, "previous_balance": prev, "new_balance": new_balance,
           "exchange_rate": rate,
           "counterpart_amount": counterpart_amt,
           "counterpart_currency": counterpart_cur,
           "description": payload.get("description", ""), "created_by_name": admin.get("name"),
           "created_at": now, "updated_at": now}
    await db.finance_transactions.insert_one(doc)
    return {"ok": True, "account": account, "previous_balance": prev, "new_balance": new_balance, "adjustment": delta}

@api_router.get("/admin/finance/monthly")
async def finance_monthly(year: Optional[int] = None, currency: str = "EGP",
                          admin: dict = Depends(require_perm("access_finance"))):
    yr = year or datetime.now(timezone.utc).year
    months = [{"month": m, "label": f"{yr}-{m:02d}", "revenue": 0.0, "cost": 0.0, "profit": 0.0} for m in range(1, 13)]
    async for t in db.finance_transactions.find({"currency": currency}):
        d = (t.get("date") or "")[:7]
        try:
            ty, tm = int(d[:4]), int(d[5:7])
        except (ValueError, IndexError):
            continue
        if ty != yr:
            continue
        typ = t.get("type"); amt = _num(t.get("amount"))
        row = months[tm - 1]
        if typ in ("income", "order_revenue"):
            row["revenue"] += amt
        elif typ == "expense":
            row["cost"] += amt
    for r in months:
        r["profit"] = r["revenue"] - r["cost"]
    return {"year": yr, "currency": currency, "months": months}

async def award_referral_points(order):
    ref = order.get("referral") or {}
    owner_id = ref.get("owner_id")
    if not owner_id:
        return
    if await db.point_transactions.find_one({"order_id": str(order["_id"]), "type": "earn"}):
        return
    per = _num(ref.get("points_per_order"))
    if per <= 0:
        return
    r = await db.referrals.find_one({"code": ref.get("code")})
    if r and _num(r.get("max_total_points")) > 0 and _num(r.get("points_awarded", 0)) + per > _num(r.get("max_total_points")):
        return
    cust = await db.customers.find_one({"_id": id_query(owner_id)})
    if not cust:
        return
    new_bal = _num(cust.get("points_available")) + per
    await db.customers.update_one({"_id": id_query(owner_id)}, {"$inc": {"points_available": per, "points_earned": per}})
    await db.point_transactions.insert_one({"customer_id": owner_id, "type": "earn", "amount": per,
        "order_id": str(order["_id"]), "reason": f"Referral order {order.get('order_number')}",
        "balance_after": new_bal, "created_at": datetime.now(timezone.utc).isoformat()})
    if r:
        await db.referrals.update_one({"_id": r["_id"]}, {"$inc": {"points_awarded": per, "reward_orders": 1}})

async def reverse_referral_points(order):
    txns = await db.point_transactions.find({"order_id": str(order["_id"]), "type": "earn"}).to_list(50)
    for t in txns:
        amt = _num(t.get("amount"))
        await db.customers.update_one({"_id": id_query(t["customer_id"])}, {"$inc": {"points_available": -amt, "points_earned": -amt}})
        await db.point_transactions.insert_one({"customer_id": t["customer_id"], "type": "reverse", "amount": -amt,
            "order_id": str(order["_id"]), "reason": f"Reversal (order dihapus/dibatalkan)", "created_at": datetime.now(timezone.utc).isoformat()})
    await db.point_transactions.delete_many({"order_id": str(order["_id"]), "type": "earn"})
    # refund redeemed points to the buyer
    redeemed = order.get("points_redeemed_le") or 0
    buyer = order.get("customer_id")
    if redeemed and buyer:
        await db.customers.update_one({"_id": id_query(buyer)}, {"$inc": {"points_available": redeemed, "points_redeemed": -redeemed}})

# ---- Customer auth ----
def gen_referral_code(username):
    base = re.sub(r"[^A-Z0-9]", "", username.upper())[:8] or "SGL"
    return f"{base}{uuid.uuid4().hex[:3].upper()}"

def clean_customer(c):
    c = clean(c)
    c.pop("password_hash", None)
    return c

@api_router.post("/customer/register")
async def customer_register(data: CustomerRegister, response: Response):
    phone = normalize_phone(data.phone)
    if not valid_intl_phone(phone):
        raise HTTPException(status_code=400, detail="No HP harus diawali + dan kode negara. Contoh: +62xxxxxxxxxx")
    if len(data.password) < 6:
        raise HTTPException(status_code=400, detail="Kata sandi minimal 6 karakter")
    if await db.customers.find_one({"phone": phone}):
        raise HTTPException(status_code=400, detail="No HP sudah terdaftar")
    if await db.customers.find_one({"username": data.username.strip()}):
        raise HTTPException(status_code=400, detail="Username sudah digunakan")
    now = datetime.now(timezone.utc).isoformat()
    doc = {"username": data.username.strip(), "phone": phone, "email": (data.email or "").strip(),
           "password_hash": hash_password(data.password), "created_at": now, "last_login": now, "active": True,
           "referral_code": gen_referral_code(data.username), "points_available": 0.0, "points_earned": 0.0, "points_redeemed": 0.0}
    r = await db.customers.insert_one(doc)
    cid = str(r.inserted_id)
    await db.referrals.insert_one({"customer_id": cid, "code": doc["referral_code"], "status": "active",
        "start_date": None, "end_date": None, "discount_percentage": 5, "max_discount_le": 100,
        "max_claim_orders": 0, "max_reward_orders": 0, "max_total_points": 0, "points_per_order": 20,
        "claims": 0, "reward_orders": 0, "points_awarded": 0, "created_at": now})
    response.set_cookie("customer_token", create_customer_token(cid), httponly=True, secure=True, samesite="none", max_age=2592000, path="/")
    return clean_customer(await db.customers.find_one({"_id": r.inserted_id}))

@api_router.post("/customer/login")
async def customer_login(data: CustomerLogin, response: Response):
    ident = data.identifier.strip()
    c = await db.customers.find_one({"$or": [{"phone": normalize_phone(ident)}, {"username": ident}]})
    if not c or not verify_password(data.password, c["password_hash"]):
        raise HTTPException(status_code=401, detail="Kredensial salah")
    if c.get("active") is False:
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan. Silakan hubungi admin.")
    await db.customers.update_one({"_id": c["_id"]}, {"$set": {"last_login": datetime.now(timezone.utc).isoformat()}})
    response.set_cookie("customer_token", create_customer_token(str(c["_id"])), httponly=True, secure=True, samesite="none", max_age=2592000, path="/")
    return clean_customer(c)

@api_router.post("/customer/logout")
async def customer_logout(response: Response):
    response.delete_cookie("customer_token", path="/")
    return {"ok": True}

@api_router.post("/customer/find-username")
async def find_username(payload: Dict[str, Any]):
    phone = normalize_phone(payload.get("phone") or "")
    if not valid_intl_phone(phone):
        raise HTTPException(status_code=400, detail="No HP harus diawali + dan kode negara. Contoh: +62xxxxxxxxxx")
    c = await db.customers.find_one({"phone": phone})
    if not c:
        raise HTTPException(status_code=404, detail="Tidak ada akun terdaftar dengan No HP ini")
    return {"username": c.get("username")}

@api_router.get("/customer/me")
async def customer_me(c: dict = Depends(get_current_customer)):
    return clean_customer(c)

@api_router.get("/customer/orders")
async def customer_orders(c: dict = Depends(get_current_customer)):
    orders = await db.orders.find({"customer_id": str(c["_id"])}).sort("created_at", -1).to_list(500)
    return [clean(o) for o in orders]

@api_router.get("/customer/points")
async def customer_points(c: dict = Depends(get_current_customer)):
    txns = await db.point_transactions.find({"customer_id": str(c["_id"])}).sort("created_at", -1).to_list(500)
    ref = await db.referrals.find_one({"customer_id": str(c["_id"])})
    return {"available": _num(c.get("points_available")), "earned": _num(c.get("points_earned")),
            "redeemed": _num(c.get("points_redeemed")), "referral_code": c.get("referral_code"),
            "referral": clean(ref) if ref else None, "transactions": [clean(t) for t in txns]}

def _guest_track_view(order):
    """Limited, non-sensitive order view for guest tracking (no address/maps/customer identifiers)."""
    items = [{"product_name_snapshot": it.get("product_name_snapshot"), "category": it.get("category"),
              "configuration_snapshot": it.get("configuration_snapshot"), "quantity": it.get("quantity"),
              "subtotal_le": it.get("subtotal_le")} for it in (order.get("items") or [order.get("item")]) if it]
    return {"order_number": order.get("order_number"), "customer_name": order.get("customer_name"),
            "items": items, "order_status": order.get("order_status"), "payment_status": order.get("payment_status"),
            "delivery_method": order.get("delivery_method"), "delivery_zone_name": order.get("delivery_zone_name"),
            "subtotal_le": order.get("subtotal_le"), "discount_le": order.get("discount_le"),
            "referral_discount_le": order.get("referral_discount_le"), "points_redeemed_le": order.get("points_redeemed_le"),
            "delivery_fee_le": order.get("delivery_fee_le"), "total_le": order.get("total_le"),
            "estimated_total_idr": order.get("estimated_total_idr"), "exchange_rate_idr_per_le": order.get("exchange_rate_idr_per_le"),
            "created_at": order.get("created_at"), "claimed": order.get("customer_id") is not None}

@api_router.post("/orders/track")
async def track_order(payload: Dict[str, Any]):
    number = (payload.get("order_number") or "").strip().upper()
    phone = normalize_phone(payload.get("phone") or "")
    if not number or not phone:
        raise HTTPException(status_code=400, detail="Nomor pesanan & No HP wajib diisi")
    order = await db.orders.find_one({"order_number": number})
    if not order or normalize_phone(order.get("customer_phone", "")) != phone:
        raise HTTPException(status_code=404, detail="Pesanan tidak ditemukan atau No HP tidak cocok")
    return _guest_track_view(order)

@api_router.post("/customer/claim-order")
async def claim_order(payload: Dict[str, Any], c: dict = Depends(get_current_customer)):
    number = (payload.get("order_number") or "").strip().upper()
    phone = normalize_phone(payload.get("phone") or "")
    if not number or not phone:
        raise HTTPException(status_code=400, detail="Nomor pesanan & No HP wajib diisi")
    order = await db.orders.find_one({"order_number": number})
    if not order or normalize_phone(order.get("customer_phone", "")) != phone:
        raise HTTPException(status_code=404, detail="Pesanan tidak ditemukan atau No HP tidak cocok")
    if order.get("customer_id"):
        if str(order["customer_id"]) == str(c["_id"]):
            raise HTTPException(status_code=400, detail="Pesanan ini sudah ada di akun kamu")
        raise HTTPException(status_code=409, detail="Pesanan sudah terhubung ke akun lain")
    # Attach account link ONLY — all historical snapshots (price/discount/referral/points/payment) stay untouched.
    await db.orders.update_one({"_id": order["_id"]}, {"$set": {"customer_id": str(c["_id"]), "customer_username": c.get("username")}})
    return {"ok": True, "order_number": order.get("order_number")}

# ---- Referral apply (public preview, requires login for real use) ----
async def validate_referral(code, subtotal, buyer_id):
    if not code:
        return 0.0, None, None
    enabled = await get_setting("referral_program_enabled", True)
    if not enabled:
        return 0.0, None, "Program referral sedang nonaktif"
    r = await db.referrals.find_one({"code": code.strip().upper()})
    if not r:
        return 0.0, None, "Kode referral tidak ditemukan"
    now = datetime.now(timezone.utc).isoformat()
    if r.get("status") != "active":
        return 0.0, None, "Kode referral tidak aktif"
    if r.get("end_date") and now > (r["end_date"] + "T23:59:59"):
        return 0.0, None, "Kode referral sudah kedaluwarsa"
    if r.get("start_date") and now < r["start_date"]:
        return 0.0, None, "Kode referral belum berlaku"
    if r.get("max_claim_orders") and _num(r.get("claims", 0)) >= _num(r["max_claim_orders"]):
        return 0.0, None, "Kuota kode referral habis"
    if buyer_id and str(r.get("customer_id")) == str(buyer_id):
        return 0.0, None, "Tidak dapat memakai kode referral sendiri"
    disc = subtotal * _num(r.get("discount_percentage")) / 100.0
    mx = _num(r.get("max_discount_le"))
    if mx > 0:
        disc = min(disc, mx)
    return round(disc, 2), r, None

@api_router.post("/referral/validate")
async def referral_validate(payload: Dict[str, Any], request: Request):
    buyer = await get_optional_customer(request)
    disc, r, err = await validate_referral(payload.get("code"), _num(payload.get("subtotal")), str(buyer["_id"]) if buyer else None)
    if err:
        return {"valid": False, "message": err, "discount_amount": 0}
    owner = await db.customers.find_one({"_id": id_query(r["customer_id"])})
    return {"valid": True, "discount_amount": disc, "percentage": _num(r.get("discount_percentage")),
            "code": r.get("code"), "owner_name": owner.get("username") if owner else ""}

# ---- Admin: referrals & customers ----
@api_router.get("/admin/referrals")
async def admin_referrals(admin: dict = Depends(require_perm("manage_settings"))):
    out = []
    for r in await db.referrals.find({}).sort("created_at", -1).to_list(500):
        r = clean(r)
        owner = await db.customers.find_one({"_id": id_query(r["customer_id"])}) if r.get("customer_id") else None
        r["owner_name"] = owner.get("username") if owner else ""
        out.append(r)
    return out

@api_router.put("/admin/referrals/{rid}")
async def admin_update_referral(rid: str, payload: Dict[str, Any], admin: dict = Depends(require_perm("manage_settings"))):
    upd = {}
    for k in ["status", "start_date", "end_date", "discount_percentage", "max_discount_le", "max_claim_orders", "max_reward_orders", "max_total_points", "points_per_order"]:
        if k in payload:
            upd[k] = payload[k]
    for k in ["discount_percentage", "max_discount_le", "max_claim_orders", "max_reward_orders", "max_total_points", "points_per_order"]:
        if k in upd:
            upd[k] = _num(upd[k])
    if "discount_percentage" in upd:
        upd["discount_percentage"] = max(0.0, min(100.0, upd["discount_percentage"]))
    await db.referrals.update_one({"_id": id_query(rid)}, {"$set": upd})
    return clean(await db.referrals.find_one({"_id": id_query(rid)}))

@api_router.get("/admin/customers")
async def admin_customers(q: Optional[str] = None, admin: dict = Depends(require_perm("manage_orders"))):
    query = {}
    if q:
        query = {"$or": [{"username": {"$regex": q, "$options": "i"}}, {"phone": {"$regex": q, "$options": "i"}}]}
    return [clean_customer(c) for c in await db.customers.find(query).sort("created_at", -1).to_list(500)]

@api_router.get("/admin/customers/{cid}")
async def admin_customer_detail(cid: str, admin: dict = Depends(require_perm("manage_orders"))):
    c = await db.customers.find_one({"_id": id_query(cid)})
    if not c:
        raise HTTPException(status_code=404, detail="Pelanggan tidak ditemukan")
    orders = await db.orders.find({"customer_id": cid}).sort("created_at", -1).to_list(500)
    txns = await db.point_transactions.find({"customer_id": cid}).sort("created_at", -1).to_list(500)
    return {"customer": clean_customer(c), "orders": [clean(o) for o in orders], "point_transactions": [clean(t) for t in txns]}

@api_router.patch("/admin/customers/{cid}")
async def admin_update_customer(cid: str, payload: Dict[str, Any], admin: dict = Depends(require_owner())):
    c = await db.customers.find_one({"_id": id_query(cid)})
    if not c:
        raise HTTPException(status_code=404, detail="Pelanggan tidak ditemukan")
    if "active" not in payload:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    await db.customers.update_one({"_id": id_query(cid)}, {"$set": {"active": bool(payload["active"])}})
    return clean_customer(await db.customers.find_one({"_id": id_query(cid)}))

@api_router.delete("/admin/customers/{cid}")
async def admin_delete_customer(cid: str, admin: dict = Depends(require_owner())):
    c = await db.customers.find_one({"_id": id_query(cid)})
    if not c:
        raise HTTPException(status_code=404, detail="Pelanggan tidak ditemukan")
    # Block permanent deletion when there is meaningful business history to preserve.
    has_orders = await db.orders.find_one({"customer_id": cid})
    referred_others = await db.orders.find_one({"referral.owner_id": cid})
    has_points_history = await db.point_transactions.find_one({"customer_id": cid})
    ref = await db.referrals.find_one({"customer_id": cid})
    ref_active_history = bool(ref and (_num(ref.get("claims", 0)) > 0 or _num(ref.get("points_awarded", 0)) > 0 or _num(ref.get("reward_orders", 0)) > 0))
    if has_orders or referred_others or has_points_history or ref_active_history or _num(c.get("points_available")) > 0 or _num(c.get("points_earned")) > 0:
        raise HTTPException(status_code=400, detail="Pelanggan memiliki riwayat bisnis (pesanan/keuangan/referral/poin). Gunakan Nonaktifkan Akun agar riwayat tetap utuh.")
    # Safe to hard-delete: no business history. Remove account + its unused referral code.
    await db.customers.delete_one({"_id": id_query(cid)})
    if ref:
        await db.referrals.delete_one({"_id": ref["_id"]})
    return {"ok": True, "deleted": True}

@api_router.post("/admin/customers/{cid}/reset-password")
async def admin_reset_customer_password(cid: str, payload: Dict[str, Any], admin: dict = Depends(require_owner())):
    c = await db.customers.find_one({"_id": id_query(cid)})
    if not c:
        raise HTTPException(status_code=404, detail="Pelanggan tidak ditemukan")
    np = payload.get("new_password") or ""
    if len(np) < 6:
        raise HTTPException(status_code=400, detail="Kata sandi minimal 6 karakter")
    await db.customers.update_one({"_id": id_query(cid)}, {"$set": {"password_hash": hash_password(np)}})
    return {"ok": True}

@api_router.post("/admin/customers/{cid}/adjust-points")
async def admin_adjust_points(cid: str, data: PointAdjustInput, admin: dict = Depends(require_owner())):
    c = await db.customers.find_one({"_id": id_query(cid)})
    if not c:
        raise HTTPException(status_code=404, detail="Pelanggan tidak ditemukan")
    new_bal = _num(c.get("points_available")) + _num(data.amount)
    if new_bal < 0:
        raise HTTPException(status_code=400, detail="Saldo poin tidak boleh negatif")
    await db.customers.update_one({"_id": id_query(cid)}, {"$inc": {"points_available": _num(data.amount)}})
    await db.point_transactions.insert_one({"customer_id": cid, "type": "adjust", "amount": _num(data.amount),
        "order_id": None, "reason": data.reason, "balance_after": new_bal, "by_name": admin.get("name"),
        "created_at": datetime.now(timezone.utc).isoformat()})
    return {"ok": True, "balance": new_bal}

# --------------------------------------------------------------------------
# Seed & migrations
# --------------------------------------------------------------------------
RAK_TYPE_B_PRICES = {
    "60_2": 565, "60_3": 755, "60_4": 985, "60_5": 1270, "60_6": 1335, "60_7": 1505,
    "80_2": 635, "80_3": 860, "80_4": 1125, "80_5": 1335, "80_6": 1545, "80_7": 1750,
    "120_2": 830, "120_3": 1180, "120_4": 1530, "120_5": 1875, "120_6": 2255, "120_7": 2570,
}

async def seed():
    admin_email = os.environ.get("ADMIN_EMAIL", "sogil.furniture@gmail.com").lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "admin123")
    existing = await db.admins.find_one({"email": admin_email})
    if not existing:
        await db.admins.insert_one({"email": admin_email, "password_hash": hash_password(admin_password),
                                    "name": "Owner Sogil", "role": "owner", "permissions": default_permissions("owner"),
                                    "created_at": datetime.now(timezone.utc).isoformat()})
    else:
        # NEVER reset/overwrite an existing Owner's password — manual changes must persist
        # across restarts/deploys. Seed only creates the Owner when none exists.
        upd = {}
        if existing.get("role") != "owner":
            upd["role"] = "owner"; upd["permissions"] = default_permissions("owner")
        if upd:
            await db.admins.update_one({"_id": existing["_id"]}, {"$set": upd})

    await db.admins.create_index("email", unique=True)
    await db.products.create_index("slug", unique=True)
    # Focused indexes to keep order/customer queries responsive as data grows.
    for spec in ["customer_id", "order_number", "order_status", "payment_status", "referral.code"]:
        await db.orders.create_index(spec)
    await db.orders.create_index([("created_at", -1)])
    await db.customers.create_index("phone")
    await db.customers.create_index("username")
    await db.point_transactions.create_index("customer_id")
    await db.point_transactions.create_index("order_id")
    await db.finance_transactions.create_index("related_order_id")
    await db.referrals.create_index("code")

    defaults = {
        "store_name": "Sogil Furniture",
        "app_name": "Sogil Furniture — Furniture Ordering & Management",
        "tagline": "Kualitas Terbaik, Untuk Ruang Terbaik",
        "store_address": "11 El-Refaey Ln, El-Darb El-Ahmar, Al-Darb Al-Ahmar, Cairo Governorate 4293042",
        "store_maps_url": "https://maps.google.com/?q=30.041889,31.262636",
        "whatsapp_number": "628XXXXXXXXXX", "exchange_rate_idr_per_le": 357,
        "bank_info": "PLACEHOLDER — Isi informasi rekening/bank melalui dashboard admin.", "logo_url": "",
        "instagram": "", "facebook": "", "tiktok": "", "email": "",
        "referral_program_enabled": True, "point_redeem_max_pct": 50,
    }
    for k, v in defaults.items():
        if await db.settings.find_one({"key": k}) is None:
            await set_setting(k, v)

    if await db.delivery_zones.count_documents({}) == 0:
        for name, fee in [("Darasah", 150), ("Gamaliyah", 200), ("Buuts", 250), ("Hay Sadis", 500),
                          ("Hay Sabi", 500), ("Hay Tsamin", 500), ("Hay Asyir", 600)]:
            await db.delivery_zones.insert_one({"name": name, "fee_le": float(fee), "active": True,
                                                "created_at": datetime.now(timezone.utc).isoformat()})

    if await db.products.count_documents({}) == 0:
        now = datetime.now(timezone.utc).isoformat()
        rak_pricing = {
            "lengths": ["60", "80", "120"], "levels": ["2", "3", "4", "5", "6", "7"],
            "base_prices": dict(RAK_TYPE_B_PRICES),
            "types": ["B", "A"],
            "type_adjustments": {"B": {"60": 0, "80": 0, "120": 0}, "A": {"60": 60, "80": 80, "120": 120}},
            "finishings": ["Natural", "Pernis", "Cat Warna"],
            "finishing": {"Natural": 0, "Pernis": {"60": 15, "80": 20, "120": 30}, "Cat Warna": {"60": 20, "80": 25, "120": 35}},
        }
        meja_pricing = {"sizes": ["40x80", "50x80"], "heights": ["30", "75"],
                        "base_prices": {"40x80_30": 700, "40x80_75": 800, "50x80_30": 800, "50x80_75": 900},
                        "finishings": ["Natural", "Pernis", "Cat Warna"], "finishing": {"Natural": 0, "Pernis": 40, "Cat Warna": 50}}
        mrv = ["Meja 40x80 + Rak 3 Tingkat", "Meja 50x80 + Rak 3 Tingkat"]
        mr_pricing = {"variants": mrv, "base_prices": {mrv[0]: 2200, mrv[1]: 2400}, "types": ["B", "A"],
                      "type_adjustments": {"B": {mrv[0]: 0, mrv[1]: 0}, "A": {mrv[0]: 80, mrv[1]: 80}},
                      "finishings": ["Natural", "Pernis", "Cat Warna"], "finishing": {"Natural": 0, "Pernis": 60, "Cat Warna": 75}}
        products = [
            {"name": "Rak Kayu", "slug": "rak-kayu", "category": "rak", "configurable": True,
             "description": "Rak serbaguna, pilih panjang, jumlah tingkat, tipe, dan finishing sesuai kebutuhanmu.",
             "starting_price_le": RAK_TYPE_B_PRICES["60_2"], "pricing": rak_pricing, "sort_order": 1},
            {"name": "Meja", "slug": "meja", "category": "meja", "configurable": True,
             "description": "Meja praktis untuk belajar atau kerja. Pilih ukuran, tinggi, dan finishing.",
             "starting_price_le": 700, "pricing": meja_pricing, "sort_order": 2},
            {"name": "Meja Rak", "slug": "meja-rak", "category": "meja_rak", "configurable": True,
             "description": "Kombinasi meja dengan rak di atasnya. Hemat ruang, multifungsi.",
             "starting_price_le": 2200, "pricing": mr_pricing, "sort_order": 3},
            {"name": "Papan Tulis", "slug": "papan-tulis", "category": "papan_tulis", "configurable": False,
             "description": "Detail pemesanan akan dikonfirmasi melalui WhatsApp.", "starting_price_le": 0, "pricing": {}, "sort_order": 4},
            {"name": "Gantungan Baju", "slug": "gantungan-baju", "category": "gantungan_baju", "configurable": False,
             "description": "Detail pemesanan akan dikonfirmasi melalui WhatsApp.", "starting_price_le": 0, "pricing": {}, "sort_order": 8},
            {"name": "Pesanan Custom", "slug": "pesanan-custom", "category": "custom", "configurable": False,
             "description": "Punya kebutuhan khusus? Ukuran custom dapat dikonsultasikan dengan admin.", "starting_price_le": 0, "pricing": {}, "sort_order": 9},
        ]
        for p in products:
            p.update({"image_url": "", "active": True, "photos": [], "representative_photo_id": None, "created_at": now, "updated_at": now})
            await db.products.insert_one(p)

    # Migration: Rak Type B base prices (idempotent via flag)
    if await db.settings.find_one({"key": "migration_rak_typeB_v2"}) is None:
        rak = await db.products.find_one({"category": "rak"})
        if rak:
            pricing = rak.get("pricing", {}) or {}
            pricing["base_prices"] = dict(RAK_TYPE_B_PRICES)
            await db.products.update_one({"_id": rak["_id"]}, {"$set": {"pricing": pricing, "starting_price_le": RAK_TYPE_B_PRICES["60_2"]}})
        await set_setting("migration_rak_typeB_v2", True)

    # Migration: Papan Tulis small sizes (30x50, 40x60, 50x70) (idempotent via flag)
    if await db.settings.find_one({"key": "migration_papan_tulis_small_sizes_v1"}) is None:
        pt = await db.products.find_one({"slug": "papan-tulis"})
        if pt:
            pricing = pt.get("pricing", {}) or {}
            groups = pricing.get("groups", []) or []
            # Ensure size options are ordered from smallest to largest
            for g in groups:
                if g.get("key") == "size":
                    current_options = g.get("options", [])
                    new_sizes = ["30x50 cm", "40x60 cm", "50x70 cm"]
                    combined = [s for s in new_sizes if s not in current_options] + current_options
                    # Sort canonically: small sizes first, then existing sizes in exact specified order
                    canonical_order = [
                        "30x50 cm", "40x60 cm", "50x70 cm",
                        "80x60 cm", "120x60 cm", "120x80 cm",
                        "180x80 cm", "180x120 cm", "240x120 cm"
                    ]
                    g["options"] = [s for s in canonical_order if s in combined]
            # Add base prices for Gantung small sizes without altering existing base prices
            base_prices = pricing.get("base_prices", {}) or {}
            base_prices["Gantung | 30x50 cm"] = 300
            base_prices["Gantung | 40x60 cm"] = 350
            base_prices["Gantung | 50x70 cm"] = 400
            pricing["groups"] = groups
            pricing["base_prices"] = base_prices
            await db.products.update_one(
                {"_id": pt["_id"]},
                {"$set": {"pricing": pricing, "starting_price_le": 300}}
            )
        await set_setting("migration_papan_tulis_small_sizes_v1", True)

    # Ensure types in pricing do not contain deprecated A+ and B+
    await db.products.update_many(
        {"pricing.types": {"$in": ["A+", "B+"]}},
        {"$pull": {"pricing.types": {"$in": ["A+", "B+"]}}}
    )

    # Ensure photos field exists on all products
    await db.products.update_many({"photos": {"$exists": False}}, {"$set": {"photos": [], "representative_photo_id": None}})

    # Seed default categories if not present
    if await db.categories.count_documents({}) == 0:
        now = datetime.now(timezone.utc).isoformat()
        default_cats = [
            {"key": "rak", "name": "Rak", "description": "Rapikan kitab, buku, dan barang dengan lebih teratur.", "sort_order": 1, "active": True, "created_at": now, "updated_at": now},
            {"key": "meja", "name": "Meja", "description": "Meja praktis untuk belajar atau kerja.", "sort_order": 2, "active": True, "created_at": now, "updated_at": now},
            {"key": "meja_rak", "name": "Meja Rak", "description": "Kombinasi meja dengan rak di atasnya. Hemat ruang, multifungsi.", "sort_order": 3, "active": True, "created_at": now, "updated_at": now},
            {"key": "papan_tulis", "name": "Papan Tulis", "description": "Cocok untuk belajar, mengajar, dan berbagai kebutuhan.", "sort_order": 4, "active": True, "created_at": now, "updated_at": now},
            {"key": "blockboard", "name": "BlackBoard", "description": "Pilihan papan tulis untuk kebutuhan belajar dan aktivitasmu.", "sort_order": 5, "active": True, "created_at": now, "updated_at": now},
            {"key": "gantungan_baju", "name": "Gantungan Baju", "description": "Gantungan baju kokoh dan hemat tempat.", "sort_order": 6, "active": True, "created_at": now, "updated_at": now},
            {"key": "custom", "name": "Koleksi Custom", "description": "Kumpulan karya dan desain custom pilihan dari Sogil Furniture.", "sort_order": 7, "active": True, "created_at": now, "updated_at": now},
        ]
        await db.categories.insert_many(default_cats)

    await db.products.update_many(
        {"category": "custom"},
        {"$set": {"name": "Koleksi Custom"}}
    )

@app.on_event("startup")
async def startup():
    await seed()

@app.on_event("shutdown")
async def shutdown():
    client.close()

# --------------------------------------------------------------------------
# Finance statistics, customer password, XLSX exports
# --------------------------------------------------------------------------
@api_router.get("/admin/finance/statistics")
async def finance_statistics(period: str = "this_month", start: Optional[str] = None, end: Optional[str] = None,
                             currency: str = "EGP", admin: dict = Depends(require_perm("access_finance"))):
    s, e, _ps, _pe = _period_range(period, start, end)
    q = {"currency": currency}
    if s or e:
        q["date"] = {}
        if s:
            q["date"]["$gte"] = s
        if e:
            q["date"]["$lte"] = e
    groups = {}
    totals = {"revenue": 0.0, "transfer_income": 0.0, "cost": 0.0, "transfer_expense": 0.0}
    team_wages_map = {}
    async for t in db.finance_transactions.find(q):
        typ = t.get("type"); amt = _num(t.get("amount")); cls = t.get("classification")
        if typ == "order_revenue":
            typ = "income"; cls = "revenue"
        if typ not in ("income", "expense"):
            continue
        cat = t.get("category", "(tanpa kategori)")
        key = (typ, cat)
        g = groups.setdefault(key, {"type": typ, "category": cat, "total": 0.0, "count": 0, "classification": cls or ("revenue" if typ == "income" else "cost")})
        g["total"] += amt; g["count"] += 1
        if typ == "income":
            totals["revenue" if cls in (None, "revenue") else "transfer_income"] += amt
        else:
            totals["cost" if cls in (None, "cost") else "transfer_expense"] += amt
            # Team wage aggregation directly from finance_transactions
            rec_id = t.get("recipient_employee_id")
            if rec_id and _is_wage_category(cat):
                tw = team_wages_map.setdefault(rec_id, {
                    "id": rec_id,
                    "name": t.get("recipient_employee_name") or "Akun Tim",
                    "total": 0.0,
                    "count": 0,
                    "transactions": []
                })
                tw["total"] += amt
                tw["count"] += 1
                tw["transactions"].append({
                    "id": str(t["_id"]),
                    "date": t.get("date"),
                    "amount": amt,
                    "currency": currency,
                    "category": cat,
                    "description": t.get("description") or "",
                    "recorded_by_name": t.get("created_by_name") or "-"
                })
    income = sorted([g for g in groups.values() if g["type"] == "income"], key=lambda x: x["total"], reverse=True)
    expense = sorted([g for g in groups.values() if g["type"] == "expense"], key=lambda x: x["total"], reverse=True)
    inc_sum = sum(g["total"] for g in income) or 1
    exp_sum = sum(g["total"] for g in expense) or 1
    for g in income:
        g["pct"] = round(g["total"] / inc_sum * 100, 1)
    for g in expense:
        g["pct"] = round(g["total"] / exp_sum * 100, 1)
    totals["operating_profit"] = totals["revenue"] - totals["cost"]

    # Enrich team_wages with latest role and name from db.admins
    admin_docs = {str(a["_id"]): a for a in await db.admins.find({}).to_list(200)}
    team_wages_list = []
    for rid, tw in team_wages_map.items():
        if tw["total"] > 0:
            adm = admin_docs.get(rid)
            if adm:
                tw["name"] = adm.get("name") or tw["name"]
                tw["role"] = adm.get("role", "employee")
            else:
                tw["role"] = "employee"
            # Sort transaction history by date descending
            tw["transactions"].sort(key=lambda x: x.get("date") or "", reverse=True)
            team_wages_list.append(tw)

    # Sort primarily by total DESC (highest wage recipient first)
    team_wages_list.sort(key=lambda x: x["total"], reverse=True)

    return {"currency": currency, "period": period, "range": {"start": s, "end": e},
            "income": income, "expense": expense, "totals": totals, "team_wages": team_wages_list}

@api_router.post("/customer/change-password")
async def customer_change_password(payload: Dict[str, Any], c: dict = Depends(get_current_customer)):
    cur = payload.get("current_password") or ""
    new = payload.get("new_password") or ""
    if not verify_password(cur, c["password_hash"]):
        raise HTTPException(status_code=400, detail="Kata sandi saat ini salah")
    if len(new) < 6:
        raise HTTPException(status_code=400, detail="Kata sandi baru minimal 6 karakter")
    await db.customers.update_one({"_id": c["_id"]}, {"$set": {"password_hash": hash_password(new)}})
    return {"ok": True}

def build_xlsx(sheets):
    from openpyxl import Workbook
    import io
    wb = Workbook(); wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet((title or "Sheet")[:31])
        ws.append(headers)
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
    if not wb.sheetnames:
        wb.create_sheet("Data")
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.getvalue()

def xlsx_response(data, filename):
    return StarletteResponse(content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})

def _date_query(period, start, end):
    s, e, _ps, _pe = _period_range(period, start, end)
    q = {}
    if s or e:
        q["date"] = {}
        if s:
            q["date"]["$gte"] = s
        if e:
            q["date"]["$lte"] = e
    return q, s, e

@api_router.get("/admin/export/finance")
async def export_finance(period: str = "this_year", start: Optional[str] = None, end: Optional[str] = None,
                         admin: dict = Depends(require_perm("access_finance"))):
    q, _s, _e = _date_query(period, start, end)
    txns = await db.finance_transactions.find(q).sort("date", -1).to_list(10000)
    headers = ["Tanggal", "ID Referensi", "Tipe", "Klasifikasi", "Kategori", "Deskripsi", "Jumlah", "Mata Uang", "Kurs (IDR/EGP)", "Nilai Ekuivalen", "Mata Uang Ekuivalen", "Dari", "Ke", "Penerima", "Dicatat oleh"]
    rows = []
    for t in txns:
        amt = _num(t.get("amount")) if t.get("type") != "transfer" else _num(t.get("from_amount"))
        rows.append([(t.get("date") or "")[:19], str(t.get("_id")), t.get("type", ""), t.get("classification", ""),
                     t.get("category", ""), t.get("description", ""), amt, t.get("currency", t.get("from_account", "")),
                     t.get("exchange_rate", ""), t.get("counterpart_amount", ""), t.get("counterpart_currency", ""),
                     t.get("from_account", ""), t.get("to_account", ""), t.get("recipient_employee_name", ""), t.get("created_by_name", "")])
    return xlsx_response(build_xlsx([("Keuangan", headers, rows)]), f"keuangan_{period}.xlsx")

@api_router.get("/admin/export/orders")
async def export_orders(period: str = "this_year", start: Optional[str] = None, end: Optional[str] = None,
                        admin: dict = Depends(require_perm("manage_orders"))):
    s, e, _ps, _pe = _period_range(period, start, end)
    q = {}
    if s or e:
        q["created_at"] = {}
        if s:
            q["created_at"]["$gte"] = s
        if e:
            q["created_at"]["$lte"] = e
    orders = await db.orders.find(q).sort("created_at", -1).to_list(10000)
    headers = ["No Pesanan", "Tanggal", "Pelanggan", "No HP", "Produk & Konfigurasi", "Qty", "Subtotal", "Diskon", "Diskon Referral", "Poin", "Ongkir", "Total", "Mata Uang", "Status Pembayaran", "Status Pesanan", "Kode Referral", "Bahasa"]
    rows = []
    for o in orders:
        items = o.get("items") or ([o.get("item")] if o.get("item") else [])
        desc = "; ".join([f"{it.get('product_name_snapshot','')} [" + ", ".join(f"{k}={v}" for k, v in (it.get('configuration_snapshot') or {}).items()) + f"] x{it.get('quantity',1)}" for it in items if it])
        qty = sum(int(it.get("quantity", 1)) for it in items if it)
        rows.append([o.get("order_number", ""), (o.get("created_at") or "")[:19], o.get("customer_name", ""), o.get("customer_phone", ""),
                     desc, qty, _num(o.get("subtotal_le")), _num(o.get("discount_le")), _num(o.get("referral_discount_le")),
                     _num(o.get("points_redeemed_le")), _num(o.get("delivery_fee_le")), _num(o.get("total_le")), "LE",
                     o.get("payment_status", ""), o.get("order_status", ""), (o.get("referral") or {}).get("code", ""), o.get("language", "")])
    return xlsx_response(build_xlsx([("Pesanan", headers, rows)]), f"pesanan_{period}.xlsx")

@api_router.get("/admin/export/config-analytics")
async def export_config_analytics(admin: dict = Depends(require_perm("manage_orders"))):
    agg = {}
    async for o in db.orders.find({}):
        items = o.get("items") or ([o.get("item")] if o.get("item") else [])
        for it in items:
            if not it:
                continue
            cfg = it.get("configuration_snapshot") or {}
            cfgs = ", ".join(f"{k}={v}" for k, v in cfg.items())
            key = (it.get("category", ""), it.get("product_name_snapshot", ""), cfgs)
            g = agg.setdefault(key, {"orders": 0, "qty": 0})
            g["orders"] += 1; g["qty"] += int(it.get("quantity", 1))
    ranked = sorted(agg.items(), key=lambda x: x[1]["qty"], reverse=True)
    headers = ["Peringkat", "Kategori", "Produk", "Konfigurasi", "Jumlah Pesanan", "Total Qty Dipesan"]
    rows = [[i + 1, k[0], k[1], k[2], v["orders"], v["qty"]] for i, (k, v) in enumerate(ranked)]
    return xlsx_response(build_xlsx([("Analitik Konfigurasi", headers, rows)]), "analitik_konfigurasi.xlsx")

@api_router.get("/admin/export/customers")
async def export_customers(admin: dict = Depends(require_perm("manage_orders"))):
    custs = await db.customers.find({}).sort("created_at", -1).to_list(10000)
    headers = ["ID", "Nama Pengguna", "Email", "No HP", "Terdaftar", "Status", "Kode Referral", "Poin Tersedia", "Poin Diperoleh"]
    rows = [[str(c.get("_id")), c.get("username", ""), c.get("email", ""), c.get("phone", ""), (c.get("created_at") or "")[:19],
             "aktif" if c.get("active", True) else "nonaktif", c.get("referral_code", ""), _num(c.get("points_available")), _num(c.get("points_earned"))] for c in custs]
    return xlsx_response(build_xlsx([("Pelanggan", headers, rows)]), "pelanggan.xlsx")

@api_router.get("/admin/export/accounts")
async def export_accounts(admin: dict = Depends(require_owner())):
    admins = await db.admins.find({}).to_list(500)
    a_headers = ["ID", "Nama", "Email", "Role", "Status", "Izin", "Dibuat", "Total Upah", "Upah Tahun Ini", "Jumlah Transaksi Upah"]
    a_rows, w_rows = [], []
    w_headers = ["Karyawan", "Tanggal", "Jumlah", "Mata Uang", "Kategori", "Deskripsi", "Dicatat oleh", "ID Transaksi"]
    now = datetime.now(timezone.utc); ykey = now.strftime("%Y")
    for a in admins:
        aid = str(a["_id"])
        perms = a.get("permissions") or {}
        perm_str = ", ".join(k for k, v in perms.items() if v)
        wages = await db.employee_wages.find({"employee_id": aid}).sort("date", -1).to_list(1000)
        tot, ytot = {}, {}
        for w in wages:
            c = w.get("currency", "EGP"); amt = _num(w.get("amount"))
            tot[c] = tot.get(c, 0) + amt
            if (w.get("date") or "")[:4] == ykey:
                ytot[c] = ytot.get(c, 0) + amt
            w_rows.append([a.get("name", ""), (w.get("date") or "")[:19], amt, c, w.get("category", ""), w.get("description", ""), w.get("recorded_by_name", ""), w.get("transaction_id", "")])
        a_rows.append([aid, a.get("name", ""), a.get("email", ""), a.get("role", ""), a.get("status", "active"), perm_str,
                       (a.get("created_at") or "")[:19], "; ".join(f"{v} {c}" for c, v in tot.items()),
                       "; ".join(f"{v} {c}" for c, v in ytot.items()), len(wages)])
    return xlsx_response(build_xlsx([("Akun", a_headers, a_rows), ("Riwayat Upah", w_headers, w_rows)]), "akun_dan_upah.xlsx")

@api_router.get("/admin/export/product-prices")
async def export_product_prices(admin: dict = Depends(require_perm("modify_products"))):
    hist = await db.price_history.find({}).sort("changed_at", -1).to_list(10000)
    headers = ["Tanggal Ubah", "Produk", "Kategori", "Harga Mulai Lama (LE)", "Harga Mulai Baru (LE)", "Diubah oleh"]
    rows = [[(h.get("changed_at") or "")[:19], h.get("product_name", ""), h.get("category", ""),
             _num(h.get("old_starting_price_le")), _num(h.get("new_starting_price_le")), h.get("changed_by_name", "")] for h in hist]
    cur_headers = ["Produk", "Kategori", "Harga Mulai (LE)", "Konfigurasi Harga (JSON)"]
    cur_rows = []
    for p in await db.products.find({}).to_list(500):
        cur_rows.append([p.get("name", ""), p.get("category", ""), _num(p.get("starting_price_le")), json.dumps(p.get("pricing") or {}, ensure_ascii=False)[:32000]])
    return xlsx_response(build_xlsx([("Riwayat Perubahan Harga", headers, rows), ("Harga Saat Ini", cur_headers, cur_rows)]), "harga_produk.xlsx")

app.include_router(api_router)

_frontend = os.environ.get("FRONTEND_URL")
_origins = [o for o in os.environ.get('CORS_ORIGINS', '').split(',') if o and o != "*"]
if _frontend and _frontend not in _origins:
    _origins.append(_frontend)
if not _origins:
    _origins = ["http://localhost:3000", "http://127.0.0.1:3000"]

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=_origins,
    allow_origin_regex=r"^https?://(localhost|127\.0\.0\.1|192\.168\.\d+\.\d+|10\.\d+\.\d+\.\d+)(:\d+)?$",
    allow_methods=["*"],
    allow_headers=["*"],
)
